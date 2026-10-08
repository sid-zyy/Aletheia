"""Data Collection Module: read test data and legacy registers from JSON, CSV, Excel and SQLite databases.

Tabular layout (CSV / Excel / database), one value per row:

    section , field            , value
    proforma, kva              , 250
    proforma, limits.oil       , 35
    temp    , hours[0][1]      , 23.1

* `section` is one of SECTIONS (below). In Excel and SQLite the sheet / table name may stand in for it.
* `field` is a path: dots for named keys, [n] for list positions.
* `value` is read as a number / true / false / null when it looks like one, otherwise as text.
  Wrap it in double quotes to force text (serial number "1098").
* An optional `series` column lets one file hold several jobs; rows for other series are skipped.

`flatten()` / `export_*()` write exactly this layout, so every job can be downloaded as a template.
"""
import csv, io, json, os, re, sqlite3, tempfile

SECTIONS = ["request", "proforma", "work", "losses", "resistance", "noload", "routine", "sc", "temp", "pressure", "ids", "other"]
MAX_BYTES = 20 * 1024 * 1024


class ImportError_(ValueError):
    """Problem with an uploaded file that the user can fix."""


# ------------------------------------------------------------ nested <-> flat
def _cell(v):
    """Value -> text for a spreadsheet cell, reversible by _parse."""
    if isinstance(v, str):
        try:
            json.loads(v)
            return json.dumps(v)  # text that looks like a number/bool: keep it text
        except ValueError:
            return v if v == v.strip() and v != "" else json.dumps(v)
    return json.dumps(v)


def _parse(s):
    if s is None: return None
    if not isinstance(s, str): return s  # Excel / SQLite already typed
    t = s.strip()
    if t == "": return ""
    try:
        return json.loads(t)
    except ValueError:
        return s if t.startswith('"') else t


def flatten(data):
    """{section: nested} -> [(section, field, value)]"""
    out = []

    def walk(sec, path, v):
        if isinstance(v, dict) and v:
            for k, x in v.items(): walk(sec, f"{path}.{k}" if path else str(k), x)
        elif isinstance(v, list) and v:
            for i, x in enumerate(v): walk(sec, f"{path}[{i}]", x)
        else:
            out.append((sec, path, v))

    for sec in SECTIONS:
        if sec in data: walk(sec, "", data[sec])
    return out


_TOKEN = re.compile(r"\[(\d+)\]|([^.\[\]]+)")


def _tokens(path):
    path = str(path).strip()
    toks = [int(i) if i != "" else k for i, k in _TOKEN.findall(path)]
    if not toks or "".join(f"[{t}]" if isinstance(t, int) else "." + t for t in toks).lstrip(".") != path:
        raise ImportError_(f"Cannot read field path '{path}' (use names.like.this and [0] for list positions)")
    return toks


def unflatten(rows, series=None):
    """[(section, field, value[, series])] -> ({section: nested}, skipped_row_count)"""
    data, skipped = {}, 0
    for n, row in enumerate(rows, 1):
        sec, field, value = (str(row[0] or "").strip().lower(), row[1], row[2])
        rs = row[3] if len(row) > 3 else None
        if not sec and (field is None or str(field).strip() == ""): continue  # blank line
        if rs not in (None, "") and series and str(rs).strip().upper() != series.upper():
            skipped += 1; continue
        if sec not in SECTIONS: raise ImportError_(f"Row {n}: unknown section '{sec}'. Expected one of: {', '.join(SECTIONS)}")
        value = _parse(value)
        if field is None or str(field).strip() == "":
            data[sec] = value; continue
        toks = _tokens(field)
        root = data.setdefault(sec, [] if isinstance(toks[0], int) else {})
        cur = root
        for i, t in enumerate(toks):
            last = i == len(toks) - 1
            nxt = value if last else ([] if isinstance(toks[i + 1], int) else {})
            if isinstance(t, int):
                if not isinstance(cur, list): raise ImportError_(f"Row {n}: '{field}' mixes list positions and names")
                if t > 5000: raise ImportError_(f"Row {n}: list position {t} is too large")
                while len(cur) <= t: cur.append(None)
                if last or cur[t] is None: cur[t] = nxt
                cur = cur[t]
            else:
                if not isinstance(cur, dict): raise ImportError_(f"Row {n}: '{field}' mixes list positions and names")
                if last or t not in cur: cur[t] = nxt
                cur = cur[t]
    if not data: raise ImportError_("No rows for this job were found in the file" + (f" ({skipped} rows belong to other series)" if skipped else ""))
    return data, skipped


# ------------------------------------------------------------ table readers
def _norm(h): return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _decode(raw):
    for enc in ("utf-8-sig", "cp1252"):
        try: return raw.decode(enc)
        except UnicodeDecodeError: pass
    raise ImportError_("Could not read the text encoding; save the CSV as UTF-8")


def read_csv(raw):
    text = _decode(raw)
    try: dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error: dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    if not rows: raise ImportError_("The CSV file is empty")
    return [("csv", rows[0], rows[1:])]


def read_xlsx(raw):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ImportError_("Could not open the workbook. Use .xlsx (save legacy .xls files as .xlsx first)") from None
    out = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
        if rows: out.append((ws.title, rows[0], rows[1:]))
    if not out: raise ImportError_("The workbook has no data")
    return out


def read_sqlite(raw):
    if not raw.startswith(b"SQLite format 3\x00"): raise ImportError_("Not a SQLite database file")
    fd, path = tempfile.mkstemp(suffix=".db")
    try:
        with os.fdopen(fd, "wb") as f: f.write(raw)
        con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)  # read-only: the source database is never modified
        try:
            out = []
            names = [r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type IN ('table','view') AND name NOT LIKE 'sqlite_%' ORDER BY name")]
            for t in names:
                cur = con.execute(f'SELECT * FROM "{t.replace(chr(34), chr(34) * 2)}" LIMIT 200000')
                out.append((t, [d[0] for d in cur.description], [list(r) for r in cur]))
            if not out: raise ImportError_("The database has no tables")
            return out
        finally:
            con.close()
    except sqlite3.Error as e:
        raise ImportError_(f"Could not read the database: {e}") from None
    finally:
        os.unlink(path)


def kind_of(filename, raw):
    ext = os.path.splitext(filename or "")[1].lower()
    if raw.startswith(b"SQLite format 3\x00") or ext in (".db", ".sqlite", ".sqlite3"): return "sqlite"
    if raw.startswith(b"PK") or ext in (".xlsx", ".xlsm"): return "xlsx"
    if ext == ".xls": raise ImportError_("Legacy .xls is not supported; save the workbook as .xlsx or .csv")
    if ext == ".json" or raw.lstrip()[:1] == b"{": return "json"
    if ext in (".csv", ".tsv", ".txt"): return "csv"
    raise ImportError_("Unsupported file type. Use .json, .csv, .xlsx or a SQLite .db file")


def read_tables(filename, raw):
    if not 0 < len(raw) <= MAX_BYTES: raise ImportError_("File must be between 1 byte and 20 MB")
    k = kind_of(filename, raw)
    return k, {"csv": read_csv, "xlsx": read_xlsx, "sqlite": read_sqlite}[k](raw)


# ------------------------------------------------------------ test data
def load_test_data(filename, raw, series=None):
    """Any supported file -> ({section: data}, kind, notes)."""
    if not 0 < len(raw) <= MAX_BYTES: raise ImportError_("File must be between 1 byte and 20 MB")
    kind = kind_of(filename, raw)
    if kind == "json":
        try: content = json.loads(_decode(raw))
        except ValueError: raise ImportError_("The JSON file is not valid JSON") from None
        return content, kind, []
    _, tables = read_tables(filename, raw)
    rows, used, ignored = [], [], []
    for name, header, body in tables:
        h = [_norm(x) for x in header]
        if "field" not in h or "value" not in h:
            ignored.append(name); continue
        fi, vi = h.index("field"), h.index("value")
        si = h.index("section") if "section" in h else None
        ri = next((h.index(a) for a in ("series", "testseries", "seriesno") if a in h), None)
        if si is None and name.strip().lower() not in SECTIONS:
            ignored.append(name); continue
        get = lambda r, i: r[i] if i is not None and i < len(r) else None
        for r in body:
            rows.append((get(r, si) if si is not None else name.strip().lower(), get(r, fi), get(r, vi), get(r, ri)))
        used.append(name)
    if not rows:
        raise ImportError_("No test data found. Expected columns: section, field, value "
                           "(or one sheet/table per section with columns field, value). Download a template from any job to see the layout.")
    data, skipped = unflatten(rows, series)
    notes = []
    if skipped: notes.append(f"{skipped} rows for other series skipped")
    if ignored and kind != "csv": notes.append("ignored: " + ", ".join(ignored))
    return data, kind, notes


# ------------------------------------------------------------ legacy register (historical records)
ALIASES = {
    "series": ("series", "testseries", "seriesno", "testseriesno", "testseriesnumber", "reportno", "testreportno", "reportnumber"),
    "sample": ("sample", "samplecode", "samplecodeno", "sampleno"),
    "customer": ("customer", "customername", "client", "manufacturer"),
    "rating": ("rating", "ratedpower", "description", "sampledescription", "equipment"),
    "address": ("address", "customeraddress"), "serial": ("serial", "serialno", "serialnumber", "slno"),
    "tests": ("tests", "testsrequested", "test", "testtype"), "criteria": ("criteria", "standard", "referencestandard"),
    "witness": ("witness", "witnessedby"), "conformity": ("conformity", "statementofconformity", "decisionrule"),
}


def load_register(filename, raw):
    """Spreadsheet / CSV / database of existing records -> (list of dicts keyed by ALIASES, source name)."""
    kind, tables = read_tables(filename, raw) if kind_of(filename, raw) != "json" else ("json", None)
    if kind == "json": raise ImportError_("Use .csv, .xlsx or a SQLite .db file for a register of records")
    for name, header, body in tables:
        h = [_norm(x) for x in header]
        col = {k: next((h.index(a) for a in al if a in h), None) for k, al in ALIASES.items()}
        if col["series"] is None or col["customer"] is None: continue
        recs = []
        for r in body:
            rec = {k: ("" if i is None or i >= len(r) or r[i] is None else str(r[i]).strip()) for k, i in col.items()}
            if any(rec.values()): recs.append(rec)
        return recs, name
    raise ImportError_("No table with at least 'series' and 'customer' columns was found. Expected columns: series, sample, customer, rating "
                       "(optional: address, serial, tests, standard, witness, conformity)")


# ------------------------------------------------------------ exporters (templates / hand-over)
HEAD = ["section", "field", "value"]


def export_csv(data):
    buf = io.StringIO(); w = csv.writer(buf, lineterminator="\n"); w.writerow(HEAD)
    for s, f, v in flatten(data): w.writerow([s, f, _cell(v)])
    return buf.getvalue().encode("utf-8-sig")


def export_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); wb.remove(wb.active)
    rows = flatten(data)
    for sec in [s for s in SECTIONS if s in data]:
        ws = wb.create_sheet(sec); ws.append(["field", "value"])
        for c in ws[1]: c.font = Font(bold=True)
        for s, f, v in rows:
            if s == sec: ws.append([f, v if isinstance(v, (int, float)) and not isinstance(v, bool) else _cell(v)])
        ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 60
    if not wb.worksheets: wb.create_sheet("request").append(["field", "value"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def export_sqlite(data, series=""):
    fd, path = tempfile.mkstemp(suffix=".db"); os.close(fd)
    try:
        con = sqlite3.connect(path)
        con.execute("CREATE TABLE test_data(series TEXT, section TEXT, field TEXT, value)")
        con.executemany("INSERT INTO test_data VALUES(?,?,?,?)",
                        [(series, s, f, v if isinstance(v, (int, float)) and not isinstance(v, bool) else _cell(v)) for s, f, v in flatten(data)])
        con.commit(); con.close()
        with open(path, "rb") as f: return f.read()
    finally:
        os.unlink(path)
