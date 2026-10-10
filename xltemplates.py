"""Excel extraction with changeable templates (docs/NEXT_STEPS.md section 5).

A template is data, not code: a JSON mapping stored in the database. When the lab changes a sheet, an administrator updates
the mapping and the Python does not change. One mapping does two jobs:

* read a filled sheet: every mapped value comes back with the cell it was read from and a status, so the tester (and later
  the tester verifying it) can see exactly where each number came from;
* write a blank sheet (and fill one from stored data), so testers always start from the current layout, with named cells.

Mapping (one sheet):

    {"section": "noload", "title": "No-Load Loss and Current Logsheet",
     "fingerprint": {"contains": "Aletheia template noload", "within": "A1:L3"},
     "fields": [
       {"field": "v100", "label": "Rated voltage V (100%)", "cell": "C5", "name": "NL_V100", "type": "number", "required": true},
       {"field": "@ids.noload[0]", "label": "Test series no.", "cell": "C3", "name": "NL_SERIES", "type": "text"},
       {"field": "rows", "type": "table", "at": "A9", "rows": 6, "required": true,
        "columns": [{"label": "Condition", "type": "text"}, {"group": ["V1", "V2", "V3"]}, {"label": "Avg V"}, ...]}]}

Where a value is looked for, in order (a field may restrict this with "by"): the workbook's defined name, the label (the
first cell after the label, skipping merged cells; "direction": "down" looks below), then the fixed cell. Labels match
after normalising case, spaces and punctuation, and fuzzily (renamed slightly) when no exact match exists; a label found
twice is reported as ambiguous instead of guessed.

Tables are found by their header labels anywhere on the sheet, so inserted rows or reordered columns do not break them.
Rows are read until the first completely blank row. Options: "rows_as": "list" (default) | "dict" (first column is the key) |
"values" (one column, a plain list) | "row" with "key" (one row picked by its first column); "pick": n returns only column n
of every row; "keys": [...] keeps only rows whose first column is listed.

Field paths are those of importers.flatten ("limits.oil", "amb[0]"); "@ids.<section>[n]" writes into the identifiers section.

Values are taken as logged (F5): a formula cell gives the value the lab's sheet shows. A formula with no saved value (the file
was never recalculated) is an error, never a guess. Dates become dd-mm-yyyy text. Decimal commas are converted only when the
template says so ("decimal": ","). Units are never guessed.
"""
import datetime as dt, difflib, io, re, zipfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

MAX_ROWS = 2000


class TemplateError(ValueError):
    """A file or mapping problem the user can fix."""


# ------------------------------------------------------------------ opening a workbook
class Book:
    """A workbook opened twice: cached values (what the sheet shows) and formulas (to spot never-recalculated cells)."""
    def __init__(self, raw, name="upload.xlsx"):
        low = name.lower()
        if low.endswith(".xls") or raw[:4] == b"\xd0\xcf\x11\xe0":
            raise TemplateError("Old .xls workbooks cannot be read; open the file in Excel and save it as .xlsx")
        if not raw.startswith(b"PK"): raise TemplateError("This is not an Excel workbook (.xlsx)")
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if any(n.lower().endswith("vbaproject.bin") for n in z.namelist()) or low.endswith(".xlsm"):
                    raise TemplateError("Macro-enabled workbooks (.xlsm) are refused; save a copy as .xlsx without macros")
            self.v = load_workbook(io.BytesIO(raw), data_only=True)
            self.f = load_workbook(io.BytesIO(raw), data_only=False)
        except TemplateError: raise
        except Exception as e:  # noqa: BLE001 - zip / XML damage
            raise TemplateError(f"The workbook could not be opened ({type(e).__name__}); it may be damaged") from None
        self.names = {}
        for n, dn in self.v.defined_names.items():
            for title, coord in dn.destinations: self.names[n.lower()] = (title, coord.replace("$", ""))
        for ws in self.v.worksheets:  # names scoped to one sheet
            for n, dn in getattr(ws, "defined_names", {}).items():
                for title, coord in dn.destinations: self.names.setdefault(n.lower(), (title or ws.title, coord.replace("$", "")))

    def sheets(self, hidden=False):
        return [ws for ws in self.v.worksheets if hidden or ws.sheet_state == "visible"]


def norm(s): return re.sub(r"[^a-z0-9%]+", " ", str(s or "").lower()).strip()


class Sheet:
    """One worksheet with merged-cell awareness and a record of which cells the mapping used."""
    def __init__(self, book, ws):
        self.b, self.ws, self.wf = book, ws, book.f[ws.title]
        self.used = set()
        self.merged = {}
        for r in ws.merged_cells.ranges:
            c1, r1, c2, r2 = r.min_col, r.min_row, r.max_col, r.max_row
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1): self.merged[(rr, cc)] = (r1, c1, r2, c2)
        self._labels = None

    def top_left(self, r, c):
        m = self.merged.get((r, c)); return (m[0], m[1]) if m else (r, c)

    def raw(self, r, c):
        r, c = self.top_left(r, c); return self.ws.cell(r, c).value

    def formula_unsaved(self, r, c):
        r, c = self.top_left(r, c)
        f = self.wf.cell(r, c).value
        return isinstance(f, str) and f.startswith("=") and self.ws.cell(r, c).value is None

    def is_formula(self, r, c):
        r, c = self.top_left(r, c); f = self.wf.cell(r, c).value
        return isinstance(f, str) and f.startswith("=")

    def ref(self, r, c): return f"{self.ws.title}!{get_column_letter(c)}{r}"

    def labels(self):
        """Text cells by normalised text -> [(row, col)]."""
        if self._labels is None:
            self._labels = {}
            for row in self.ws.iter_rows(max_row=min(self.ws.max_row, MAX_ROWS)):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.strip():
                        self._labels.setdefault(norm(cell.value), []).append((cell.row, cell.column))
        return self._labels

    def find_label(self, label, near=None):
        """(row, col) of the label, or raises LookupError('missing' | 'ambiguous') with detail."""
        L = self.labels(); key = norm(label)
        hits = L.get(key)
        if not hits:  # tolerate small renames ("Rated volt. (100%)" for "Rated voltage (100%)")
            best = difflib.get_close_matches(key, list(L), n=2, cutoff=0.86)
            if best: hits = L[best[0]]
            if len(best) == 2 and difflib.SequenceMatcher(None, key, best[1]).ratio() == difflib.SequenceMatcher(None, key, best[0]).ratio():
                raise LookupError("ambiguous", f"'{label}' is close to both '{best[0]}' and '{best[1]}'")
        if not hits: raise LookupError("missing", f"label '{label}' not found")
        if len(hits) > 1:
            if near:  # prefer the occurrence nearest the template's own cell
                hits = sorted(hits, key=lambda h: abs(h[0] - near[0]) + abs(h[1] - near[1]))
                if abs(hits[0][0] - near[0]) + abs(hits[0][1] - near[1]) < abs(hits[1][0] - near[0]) + abs(hits[1][1] - near[1]): return hits[0]
            raise LookupError("ambiguous", f"label '{label}' appears {len(hits)} times ({', '.join(self.ref(*h) for h in hits[:4])})")
        return hits[0]

    def after(self, r, c, direction="right"):
        """The value cell next to a label: past the label's merged area."""
        m = self.merged.get((r, c))
        if direction == "down": return ((m[2] if m else r) + 1, c)
        return (r, (m[3] if m else c) + 1)


def addr(cell):
    """'Sheet!C5' or 'C5' -> (sheet or None, row, col)."""
    sh, _, a = cell.rpartition("!")
    m = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?(\d+)", a.strip())
    if not m: raise TemplateError(f"Not a cell reference: {cell}")
    return (sh.strip("'") or None), int(m.group(2)), column_index_from_string(m.group(1).upper())


# ------------------------------------------------------------------ values
def convert(v, typ, opts, sheet=None, rc=None):
    """Cell value -> (value, problem). Never invents a value: anything it cannot read as the declared type is a problem."""
    if isinstance(v, str): v = v.strip()
    if v is None or v == "": return (opts.get("blank", None), None)
    if typ in (None, "any"):
        if isinstance(v, (dt.datetime, dt.date)): return v.strftime("%d-%m-%Y"), None
        return v, None
    if typ == "number":
        if isinstance(v, bool): return None, f"expected a number, found {v}"
        if isinstance(v, (int, float)):
            if sheet is not None and opts.get("scale"): v = v * opts["scale"]
            return (int(v) if isinstance(v, float) and v.is_integer() and opts.get("int") else v), None
        s = str(v).replace("−", "-").replace(" ", "")
        if opts.get("decimal") == ",": s = s.replace(".", "").replace(",", ".")
        if re.fullmatch(r"[-+]?\d+(\.\d+)?([eE][-+]?\d+)?", s):
            x = float(s); return (int(x) if x.is_integer() and "." not in s and "e" not in s.lower() else x), None
        return None, f"expected a number, found '{v}'" + (" (a decimal comma? the template does not allow it)" if re.fullmatch(r"[-+]?\d+,\d+", s) else "")
    if typ == "text":
        if isinstance(v, float) and v.is_integer(): v = int(v)
        if isinstance(v, (dt.datetime, dt.date)): return v.strftime("%d-%m-%Y"), None
        return str(v), None
    if typ == "date":
        if isinstance(v, (dt.datetime, dt.date)): return v.strftime("%d-%m-%Y"), None
        s = str(v)
        for f in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d"):
            try: return dt.datetime.strptime(s, f).strftime("%d-%m-%Y"), None
            except ValueError: pass
        return s, None  # free-text dates ("09-11-2025 to 10-11-2025") are kept as written
    if typ == "bool":
        t = str(v).strip().lower()
        if t in ("yes", "y", "true", "1"): return True, None
        if t in ("no", "n", "false", "0"): return False, None
        return None, f"expected yes or no, found '{v}'"
    raise TemplateError(f"Unknown type '{typ}' in the template")


# ------------------------------------------------------------------ paths
_TOK = re.compile(r"\[(\d+)\]|([^.\[\]]+)")


def tokens(path):
    t = [int(i) if i != "" else k for i, k in _TOK.findall(str(path))]
    if not t: raise TemplateError(f"Empty field path in the template")
    return t


def put(root, path, value):
    toks = tokens(path); cur = root
    for i, t in enumerate(toks):
        last = i == len(toks) - 1
        nxt = value if last else ([] if isinstance(toks[i + 1], int) else {})
        if isinstance(t, int):
            if not isinstance(cur, list): raise TemplateError(f"Field path '{path}' mixes list positions and names")
            while len(cur) <= t: cur.append(None)
            if last or cur[t] is None: cur[t] = nxt
            cur = cur[t]
        else:
            if not isinstance(cur, dict): raise TemplateError(f"Field path '{path}' mixes list positions and names")
            if last or t not in cur: cur[t] = nxt
            cur = cur[t]


def get(root, path):
    cur = root
    for t in tokens(path):
        try: cur = cur[t]
        except (KeyError, IndexError, TypeError): return None
    return cur


# ------------------------------------------------------------------ fingerprint / detection
def matches(sheet, fp):
    """Does this worksheet carry the template's fingerprint?"""
    if not fp: return False
    if fp.get("sheet") and norm(fp["sheet"]) != norm(sheet.ws.title): return False
    if "contains" in fp or "equals" in fp:
        c1, r1, c2, r2 = range_boundaries(fp.get("within") or "A1:Z10")
        want = norm(fp.get("contains") or fp.get("equals"))
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                t = norm(sheet.ws.cell(r, c).value) if isinstance(sheet.ws.cell(r, c).value, str) else ""
                if t and ((want in t) if "contains" in fp else t == want): return True
        return False
    if "label" in fp:
        try: sheet.find_label(fp["label"]); return True
        except LookupError: return False
    return bool(fp.get("sheet"))


def detect(book, templates):
    """[(worksheet, template)] for every visible sheet that one of the templates recognises (first match wins)."""
    out = []
    for ws in book.sheets():
        sh = Sheet(book, ws)
        t = next((t for t in templates if matches(sh, t["mapping"].get("fingerprint"))), None)
        if t: out.append((ws, t))
    return out


# ------------------------------------------------------------------ extraction
def extract(book, mapping, ws=None):
    """Read one sheet with one mapping. Returns dict(section, data, ids, fields=[{field, value, cell, status, note}], errors,
    warnings, unmapped). Nothing is stored here; the caller shows the result and saves only on confirmation."""
    section = mapping.get("section") or "other"
    if ws is None:
        hits = [s for s in book.sheets() if matches(Sheet(book, s), mapping.get("fingerprint"))]
        ws = hits[0] if hits else (book.sheets() or [None])[0]
    if ws is None: raise TemplateError("The workbook has no visible sheet")
    sh = Sheet(book, ws)
    for it in mapping.get("layout") or []:  # printed text of a paper layout is the form itself, never a value outside the mapping
        try:
            c1, r1, c2, r2 = range_boundaries(f"{it['at']}:{it.get('to') or it['at']}")
            sh.used.update((r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1))
        except (KeyError, ValueError, TypeError): pass
    data, ids, fields, errors, warnings = {}, {}, [], [], []
    if ws.sheet_state != "visible": warnings.append(f"Sheet '{ws.title}' is hidden")
    for f in mapping.get("fields") or []:
        if f.get("type") == "table" and f.get("caption") and f.get("at"):  # the table's printed caption, above its header
            try: _, cr, cc = addr(f["at"]); sh.used.add(sh.top_left(cr - 1, cc))
            except TemplateError: pass
        if f.get("type") == "table": res = read_table(sh, f, book)
        elif f.get("by") == "const": res = dict(field=f["field"], value=f.get("value"), cell=None, status="ok", note="fixed in the template")
        else: res = read_field(sh, f, book)
        if f.get("title") or f.get("label") or f.get("caption"): res["label"] = f.get("title") or f.get("label") or f.get("caption")
        fields.append(res)
        st = res["status"]
        if st in ("missing", "ambiguous", "type", "formula_unsaved", "not_found") and (f.get("required") or st != "missing"):
            errors.append(f"{f.get('title') or f.get('label') or f.get('caption') or f['field']}: {res['note']}" + (f" ({res['cell']})" if res.get("cell") else ""))
        if st in ("ok", "formula", "missing", "hidden") and not (st == "missing" and f.get("required")):
            target = f["field"]
            if target.startswith("@ids."):
                if res["value"] is not None: put(ids, target[5:], res["value"])  # an identifier not written on the sheet is left out, not stored as NA
            else: put(data, target, res["value"])
        if st == "hidden": warnings.append(f"{f.get('label') or f['field']}: read from a hidden row or column ({res['cell']})")
    data = {k: prune(v) for k, v in data.items()}
    unmapped = []
    for row in ws.iter_rows(max_row=min(ws.max_row, MAX_ROWS)):
        for cell in row:
            v = cell.value
            if v is None or (isinstance(v, str) and not v.strip()): continue
            if (cell.row, cell.column) in sh.used or isinstance(v, str) and (len(v) > 25 or not re.search(r"\d", v)): continue  # labels and notes
            unmapped.append(dict(cell=sh.ref(cell.row, cell.column), value=v if not isinstance(v, (dt.date, dt.datetime)) else v.isoformat()))
    if unmapped: warnings.append(f"{len(unmapped)} value cell{'s' if len(unmapped) > 1 else ''} outside the mapping were not read (first: {unmapped[0]['cell']})")
    return dict(section=section, sheet=ws.title, data=data, ids=ids, fields=fields, errors=errors, warnings=warnings, unmapped=unmapped[:50])


def prune(v):
    """Inside a group of values (a dict), what was left empty is left out, and a group left entirely empty is None; lists keep
    their positions (Before / After pairs). Top-level fields keep None: an empty cell is recorded as NA."""
    if not isinstance(v, dict): return v
    out = {k: prune(x) for k, x in v.items()}
    out = {k: x for k, x in out.items() if x is not None}
    return out or None


def locate(sh, f, book):
    """(row, col, how) of a single field's value cell, or raises LookupError. A name or fixed cell whose neighbour no
    longer reads as the field's label has drifted (rows inserted by a tool that does not move names): the label wins."""
    r, c, how = _locate(sh, f, book)
    if how in ("name", "cell") and f.get("label") and "label" in (f.get("by") or ["name", "label", "cell"]):
        left = sh.raw(r, c - 1) if c > 1 else None; up = sh.raw(r - 1, c) if r > 1 else None
        if norm(f["label"]) not in (norm(left), norm(up)):
            try:
                lr, lc = sh.find_label(f["label"]); sh.used.add((lr, lc))
                rr, cc = sh.after(lr, lc, f.get("direction", "right"))
                return rr, cc, "label (the " + how + " has moved away from its label)"
            except LookupError: pass
    return r, c, how


def _locate(sh, f, book):
    order = f.get("by") or ["name", "label", "cell"]
    order = [order] if isinstance(order, str) else order
    notes = []
    for how in order:
        if how == "name" and f.get("name"):
            hit = book.names.get(f["name"].lower())
            if hit:
                title, coord = hit
                if title and title != sh.ws.title: notes.append(f"name {f['name']} points to sheet '{title}'"); continue
                _, r, c = addr(coord.split(":")[0]); return r, c, "name"
            notes.append(f"name {f['name']} not defined")
        elif how == "label" and f.get("label"):
            near = addr(f["cell"])[1:] if f.get("cell") else None
            try:
                r, c = sh.find_label(f["label"], near=near); sh.used.add((r, c))
                rr, cc = sh.after(r, c, f.get("direction", "right"))
                if f.get("offset"): cc += int(f["offset"])
                return rr, cc, "label"
            except LookupError as e:
                if e.args[0] == "ambiguous": raise
                notes.append(e.args[1])
        elif how == "cell" and f.get("cell"):
            s, r, c = addr(f["cell"])
            if s and s != sh.ws.title: notes.append(f"cell {f['cell']} is on another sheet"); continue
            return r, c, "cell"
    raise LookupError("not_found", "; ".join(notes) or "no way to locate this value (give a name, label or cell)")


def read_field(sh, f, book):
    try: r, c, how = locate(sh, f, book)
    except LookupError as e:
        return dict(field=f["field"], value=None, cell=None, status="ambiguous" if e.args[0] == "ambiguous" else "not_found", note=e.args[1])
    sh.used.add(sh.top_left(r, c))
    if f.get("label") and c > 1 and norm(sh.raw(r, c - 1)) == norm(f["label"]): sh.used.add(sh.top_left(r, c - 1))  # its label
    ref = sh.ref(*sh.top_left(r, c))
    if sh.formula_unsaved(r, c):
        return dict(field=f["field"], value=None, cell=ref, status="formula_unsaved",
                    note="formula without a saved value: open the file in Excel, let it recalculate, save, and upload again")
    val, prob = convert(sh.raw(r, c), f.get("type"), f, sheet=sh, rc=(r, c))
    if prob: return dict(field=f["field"], value=None, cell=ref, status="type", note=prob, found=str(sh.raw(r, c)))
    hidden = sh.ws.row_dimensions[r].hidden or sh.ws.column_dimensions[get_column_letter(c)].hidden
    st = "missing" if val in (None, "") and sh.raw(r, c) in (None, "") else "hidden" if hidden else "formula" if sh.is_formula(r, c) else "ok"
    note = {"missing": "empty" + (" (required)" if f.get("required") else ", recorded as NA"), "formula": "formula: the value the sheet shows is taken as logged",
            "hidden": "hidden row or column"}.get(st, f"found by {how}")
    return dict(field=f["field"], value=val, cell=ref, status=st, note=note, by=how)


def col_specs(f):
    """Flatten the column specs: [(position in output, index in group or None, label, type, opts)]."""
    out = []
    for i, c in enumerate(f["columns"]):
        if "group" in c:
            for j, lab in enumerate(c["group"]): out.append((i, j, lab, c.get("type", f.get("cell_type", "number")), c))
        else: out.append((i, None, c["label"], c.get("type", f.get("cell_type", "number")), c))
    return out


def find_header(sh, f, specs):
    """Row and column of each header label. Headers may be split over two rows (a group label above V1, V2, V3)."""
    first = specs[0][2]
    try: r0, c0 = sh.find_label(first, near=addr(f["at"])[1:] if f.get("at") else None)
    except LookupError as e:
        if f.get("at"):
            _, r0, c0 = addr(f["at"])
            if norm(sh.raw(r0, c0)) != norm(first): raise LookupError(e.args[0], f"table header '{first}' not found") from None
        else: raise LookupError(e.args[0], f"table header '{first}': {e.args[1]}") from None
    cols = {}
    row_labels = {}
    for cc in range(1, sh.ws.max_column + 1):
        for rr in (r0, r0 + 1):
            v = sh.raw(rr, cc)
            if isinstance(v, str) and v.strip(): row_labels.setdefault(norm(v), []).append((rr, cc))
    hdr_bottom = r0
    for pos, j, lab, typ, c in specs:
        hits = row_labels.get(norm(lab))
        if not hits:
            best = difflib.get_close_matches(norm(lab), list(row_labels), n=1, cutoff=0.86)
            hits = row_labels.get(best[0]) if best else None
        if not hits: raise LookupError("missing", f"column '{lab}' not found next to header '{first}' ({sh.ref(r0, c0)})")
        hits = [h for h in hits if h[1] >= c0] or hits
        rr, cc = hits[0] if len(hits) == 1 else min(hits, key=lambda h: (h[1] - c0) if (h[1] not in cols.values()) else 10 ** 6)
        cols[(pos, j)] = cc; hdr_bottom = max(hdr_bottom, rr)
        sh.used.add((rr, cc))
    sh.used.add((r0, c0))
    return hdr_bottom, cols


def blank(v): return v is None or (isinstance(v, str) and not v.strip())


def is_calc(c):
    """A calculated column: drawn with a formula; read only when "read" is set (it then holds a logged value)."""
    return bool(c.get("calc")) and not c.get("read")


def read_table(sh, f, book):
    if f.get("orient") == "columns": return read_tgrid(sh, f, book)
    specs = col_specs(f)
    try: hdr, cols = find_header(sh, f, specs)
    except LookupError as e:
        return dict(field=f["field"], value=None, cell=None, status="ambiguous" if e.args[0] == "ambiguous" else "not_found", note=e.args[1])
    labels = f.get("row_labels")
    rows, problems, first_ref, r = [], [], None, hdr + 1
    last = hdr + len(labels) if labels else min(sh.ws.max_row, hdr + MAX_ROWS)
    while r <= last:
        cells = {k: (r, cc) for k, cc in cols.items()}
        raw = {k: sh.raw(*rc) for k, rc in cells.items()}
        inputs = [raw[(i, j)] for i, j in raw if not (i == 0 and labels) and not is_calc(f["columns"][i])]
        if all(blank(v) for v in inputs):
            if labels:  # a printed row label with nothing logged against it (90 %, hour 15...): skipped, not the end
                for rc in cells.values(): sh.used.add(rc)
                r += 1; continue
            break
        if f.get("stop_label") and norm(raw.get((0, None))) == norm(f["stop_label"]): break
        out = []
        for i, c in enumerate(f["columns"]):
            def one(key, typ, opts):
                rc = cells[key]; sh.used.add(rc)
                if is_calc(c): return None
                if sh.formula_unsaved(*rc): problems.append(f"{sh.ref(*rc)}: formula without a saved value")
                v, p = convert(raw[key], typ, opts)
                if p: problems.append(f"{sh.ref(*rc)}: {p}")
                if c.get("calc") and isinstance(v, float): v = round(v, 6)  # the sheet calculated it: no binary noise (382.91999999999996)
                return v
            if "group" in c: out.append([one((i, j), c.get("type", f.get("cell_type", "number")), c) for j in range(len(c["group"]))])
            else: out.append(one((i, None), c.get("type", f.get("cell_type", "number")), c))
        first_ref = first_ref or sh.ref(r, min(cols.values()))
        rows.append(out); r += 1
    ref = f"{first_ref}..{sh.ref(r - 1, max(cols.values()))}" if rows else sh.ref(hdr, min(cols.values()))
    return table_result(f, rows, problems, ref)


def read_tgrid(sh, f, book):
    """A table drawn sideways, as on several paper logsheets: the column titles run down the first column and every record
    is one column to the right of them (reference points 1 to 6, readings over time)."""
    titles = [c["label"] for c in f["columns"]]
    try: r0, c0 = sh.find_label(titles[0], near=addr(f["at"])[1:] if f.get("at") else None)
    except LookupError as e:
        return dict(field=f["field"], value=None, cell=None, status="ambiguous" if e.args[0] == "ambiguous" else "not_found", note=f"table title '{titles[0]}': {e.args[1]}")
    trow = {0: r0}
    for i, t in enumerate(titles[1:], 1):  # each title in the same column, below the previous one (rows may have been inserted)
        hit = next((rr for rr in range(trow[i - 1] + 1, trow[i - 1] + 6) if norm(sh.raw(rr, c0)) == norm(t)), None)
        if hit is None: return dict(field=f["field"], value=None, cell=sh.ref(r0, c0), status="not_found", note=f"row title '{t}' not found under '{titles[0]}'")
        trow[i] = hit
    for rr in trow.values(): sh.used.add((rr, c0))
    labels = f.get("row_labels"); rows, problems, refs = [], [], []
    for k in range(f.get("records") or len(labels or []) or 6):
        cc = c0 + 1 + k
        raw = {i: sh.raw(trow[i], cc) for i in trow}
        for i in trow: sh.used.add((trow[i], cc))
        if all(blank(raw[i]) for i, c in enumerate(f["columns"]) if not (i == 0 and labels) and not is_calc(c)): continue
        out = []
        for i, c in enumerate(f["columns"]):
            if is_calc(c): out.append(None); continue
            if sh.formula_unsaved(trow[i], cc): problems.append(f"{sh.ref(trow[i], cc)}: formula without a saved value")
            v, p = convert(raw[i], c.get("type", f.get("cell_type", "number")), c)
            if p: problems.append(f"{sh.ref(trow[i], cc)}: {p}")
            out.append(v)
        rows.append(out); refs.append(cc)
    ref = f"{sh.ref(r0, refs[0])}..{sh.ref(trow[len(titles) - 1], refs[-1])}" if refs else sh.ref(r0, c0)
    return table_result(f, rows, problems, ref)


def take_row(f, x):
    """The stored row from a sheet row: "take" lists the sheet columns (n, or [n, k] for item k of a group) in stored order."""
    if not f.get("take"): return x
    return [x[t[0]][t[1]] if isinstance(t, list) and isinstance(x[t[0]], list) else x[t] if not isinstance(t, list) else None for t in f["take"]]


def table_result(f, rows, problems, ref):
    if problems:
        return dict(field=f["field"], value=None, cell=ref, status="type", note="; ".join(problems[:5]) + (f" (+{len(problems) - 5} more)" if len(problems) > 5 else ""))
    rows = [take_row(f, x) for x in rows]
    if f.get("take"):  # this field reads part of a shared table: a row with nothing in its part is not one of its rows
        rows = [x for x in rows if any(not blank(v) and not (isinstance(v, list) and all(blank(y) for y in v)) for v in (x[1:] if f["take"][0] == 0 else x))]
    if not rows:
        return dict(field=f["field"], value=None, cell=ref, status="missing", note="table is empty" + (" (required)" if f.get("required") else ""))
    if f.get("keys_as"):  # printed row labels stored under short keys ("Normal tap" -> "NT")
        km = {norm(k): v for k, v in f["keys_as"].items()}
        rows = [[km.get(norm(x[0]), x[0])] + list(x[1:]) for x in rows]
    if f.get("keys"): rows = [x for x in rows if x[0] in f["keys"]]
    how = f.get("rows_as", "list")
    if "pick" in f: val = [x[f["pick"]] for x in rows]
    elif how == "dict": val = {str(x[0]): (x[1] if len(x) == 2 else x[1:]) for x in rows}
    elif how == "values": val = [x[0] for x in rows]
    elif how == "first": val = rows[0]
    elif how == "row":
        hit = [x for x in rows if str(x[0]) == str(f.get("key"))]
        if not hit: return dict(field=f["field"], value=None, cell=ref, status="missing", note=f"no row '{f.get('key')}' in the table")
        val = hit[0][1] if len(hit[0]) == 2 else hit[0][1:]
    else: val = rows
    return dict(field=f["field"], value=val, cell=ref, status="ok", note=f"{len(rows)} row{'s' if len(rows) != 1 else ''}")


# ------------------------------------------------------------------ writing: blank sheets and filled sheets
THIN = Side(style="thin", color="9AA8BD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="DFE7F1")
INPUT = PatternFill("solid", fgColor="FFFBEA")


def draw(ws, mapping, data=None, ids=None, wb=None):
    """Lay out one sheet from its mapping (labels, input cells, named cells, table headers); fill it from data if given.
    A mapping with a "layout" (version 2 and later) is drawn as the paper logsheet instead: see draw_paper."""
    if mapping.get("layout") is not None: return draw_paper(ws, mapping, data, ids, wb)
    ws["A1"] = mapping.get("title") or mapping.get("section"); ws["A1"].font = Font(bold=True, size=13)
    fp = mapping.get("fingerprint") or {}
    if fp.get("contains"):
        ws["A2"] = f"CPRI Short Circuit Laboratory - form: {fp['contains']} (version {mapping.get('version', 1)}). Fill the yellow cells only; do not move or rename them."
        ws["A2"].font = Font(italic=True, size=9, color="56667F")
    ws.column_dimensions["A"].width = 30
    tables = {}
    for f in mapping.get("fields") or []:
        if f.get("by") == "const": continue
        if f.get("type") == "table": tables.setdefault(f["at"], []).append(f); continue
        _, r, c = addr(f["cell"])
        if f.get("label") and not ws.cell(r, c - 1).value:
            lc = ws.cell(r, c - 1, f["label"]); lc.alignment = Alignment(wrap_text=True, vertical="top")
        vc = ws.cell(r, c); vc.fill = INPUT; vc.border = BOX
        w = ws.column_dimensions[get_column_letter(c)].width or 0
        ws.column_dimensions[get_column_letter(c)].width = max(w, 24 if f.get("type") == "text" else 14)
        if data is not None or ids is not None:
            v = get(ids or {}, f["field"][5:]) if f["field"].startswith("@ids.") else get(data or {}, f["field"])
            if v is not None: vc.value = v if not isinstance(v, (list, dict)) else str(v)
            if f.get("type") == "text" and isinstance(vc.value, str): vc.number_format = "@"
        if f.get("name") and wb is not None:
            wb.defined_names[f["name"]] = DefinedName(f["name"], attr_text=f"'{ws.title}'!${get_column_letter(c)}${r}")
    for at, fs in tables.items(): draw_table(ws, at, fs, data)


def draw_table(ws, at, fs, data):
    """One table region (possibly read by several fields): headers once, then the rows of all its fields."""
    f0 = fs[0]; _, r0, c0 = addr(at); specs = col_specs(f0)
    two = any(cs.get("title") for *_, cs in specs)
    if f0.get("caption"): ws.cell(r0 - 1, c0, f0["caption"]).font = Font(bold=True)
    c, done = c0, set()
    for pos, j, lab, typ, cs in specs:
        hr = r0 + 1 if (two and j is not None) else r0
        cell = ws.cell(hr, c, lab); cell.font = Font(bold=True); cell.fill = HEAD; cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        if two and j is None: ws.merge_cells(start_row=r0, start_column=c, end_row=r0 + 1, end_column=c)
        if two and j == 0 and pos not in done:
            n = len(f0["columns"][pos]["group"]); done.add(pos)
            t = ws.cell(r0, c, cs.get("title") or ""); t.font = Font(bold=True); t.fill = HEAD; t.alignment = Alignment(horizontal="center")
            if n > 1: ws.merge_cells(start_row=r0, start_column=c, end_row=r0, end_column=c + n - 1)
        w = ws.column_dimensions[get_column_letter(c)].width or 0
        ws.column_dimensions[get_column_letter(c)].width = max(w, 12)
        c += 1
    first = r0 + (2 if two else 1)
    rows = []
    if data is not None:
        if any("pick" in f for f in fs):
            merged = {}
            for f in fs:
                for i, x in enumerate(get(data, f["field"]) or []): merged.setdefault(i, [None] * len(f0["columns"]))[f["pick"]] = x
            rows = [merged[i] for i in sorted(merged)]
        else:
            for f in fs: rows += table_rows(f, get(data, f["field"]))
    for i in range(max(max(f.get("rows", 5) for f in fs), len(rows))):
        out = rows[i] if i < len(rows) else None
        cc = c0
        for pos, j, lab, typ, cs in specs:
            cell = ws.cell(first + i, cc); cell.border = BOX; cell.fill = INPUT
            if out is not None:
                v = out[pos] if j is None else (out[pos][j] if isinstance(out[pos], list) and j < len(out[pos]) else None)
                if v is not None: cell.value = v
            cc += 1


def table_rows(f, val):
    """Stored value -> rows in the table's column order (the inverse of read_table)."""
    if val is None: return []
    how, wide = f.get("rows_as", "list"), len(f["columns"]) > 2
    if how == "dict": return [[k] + (list(v) if wide else [v]) for k, v in val.items()]
    if how == "values": return [[x] for x in val]
    if how == "row": return [[f.get("key")] + (list(val) if wide else [val])]
    return [list(x) for x in val]


def workbook(sheets):
    """[(mapping, data or None, ids or None)] -> .xlsx bytes, one sheet per mapping (blank when data is None)."""
    wb = Workbook(); wb.remove(wb.active)
    for mapping, data, ids in sheets:
        title = re.sub(r"[\\/*?:\[\]]", "", mapping.get("sheet_title") or mapping.get("title") or mapping.get("section") or "Sheet")[:31]
        n, base = 2, title
        while title in wb.sheetnames: title = f"{base[:28]} {n}"; n += 1
        draw(wb.create_sheet(title), mapping, data, ids, wb)
    out = io.BytesIO(); wb.save(out); return out.getvalue()


# ---- paper layout (template version 2): the sheet drawn as the laboratory's printed logsheet
FIXED = PatternFill("solid", fgColor="EEF1F5")   # printed text and row labels: locked
CALC = PatternFill("solid", fgColor="E3E6EA")    # calculated by the sheet: locked, grey
NOTE = PatternFill("solid", fgColor="EAF3FB")
STYLE = {
    "org": dict(font=Font(bold=True, size=12), align=Alignment(horizontal="center", vertical="center")),
    "unit": dict(font=Font(size=11), align=Alignment(horizontal="center", vertical="center")),
    "title": dict(font=Font(bold=True, size=13, underline="single"), align=Alignment(horizontal="center", vertical="center")),
    "meta": dict(font=Font(size=9, color="333333"), align=Alignment(vertical="center", wrap_text=True)),
    "meta_r": dict(font=Font(size=9, color="333333"), align=Alignment(horizontal="right", vertical="center", wrap_text=True)),
    "note": dict(font=Font(italic=True, size=9, color="1C3D6E"), align=Alignment(vertical="top", wrap_text=True), fill=NOTE),
    "section": dict(font=Font(bold=True, size=11), align=Alignment(horizontal="center", vertical="center"), fill=HEAD, border=True),
    "head": dict(font=Font(bold=True), align=Alignment(horizontal="center", vertical="center", wrap_text=True), fill=HEAD, border=True),
    "label": dict(font=Font(size=10), align=Alignment(vertical="center", wrap_text=True), fill=FIXED, border=True),
    "label_b": dict(font=Font(bold=True, size=10), align=Alignment(vertical="center", wrap_text=True), fill=FIXED, border=True),
    "text": dict(font=Font(size=9), align=Alignment(vertical="top", wrap_text=True), border=True),
    "sig": dict(font=Font(bold=True, size=10), align=Alignment(horizontal="center", vertical="top")),
    "foot": dict(font=Font(size=8, color="56667F"), align=Alignment(vertical="center", wrap_text=True)),
}
DATE_MIN = 36526  # 01-01-2000 as an Excel serial: dates before it are typing mistakes


def rng_cells(ws, a, b=None):
    c1, r1, c2, r2 = range_boundaries(f"{a}:{b or a}")
    return [ws.cell(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def place(ws, at, to=None, value=None, style=None, fill=None, locked=True):
    """One cell or merged area: value in the top-left, the style and border on every cell of the area."""
    cells = rng_cells(ws, at, to)
    st = STYLE.get(style, {})
    for c in cells:  # styled before merging: Excel draws a merged area's border from every cell of it
        if st.get("border") or style is None: c.border = BOX
        if st.get("fill") or fill: c.fill = fill or st["fill"]
        c.protection = Protection(locked=locked)
    top = cells[0]
    if value is not None: top.value = value
    if st.get("font"): top.font = st["font"]
    top.alignment = st.get("align") or Alignment(horizontal="left", vertical="center", wrap_text=True)
    if to and to != at and f"{at}:{to}" not in {str(r) for r in ws.merged_cells.ranges}: ws.merge_cells(f"{at}:{to}")
    return top


def validate_cells(ws, ref, spec, typ):
    """Stop mistakes at entry: a list of choices, a number within limits, or a date."""
    title = spec.get("title") or spec.get("label") or ""
    if spec.get("choices"):
        dv = DataValidation(type="list", formula1='"' + ",".join(str(x) for x in spec["choices"]) + '"', allow_blank=True)
        dv.error = "Choose one of the listed values" + (" (or keep what you typed if it is on the paper sheet)" if spec.get("other") else "")
        dv.errorStyle = "warning" if spec.get("other") else "stop"
    elif typ == "number":
        dv = DataValidation(type="decimal", operator="between", formula1=str(spec.get("min", -1e9)), formula2=str(spec.get("max", 1e9)), allow_blank=True)
        dv.error = f"Enter a number" + (f" between {spec['min']} and {spec['max']}" if "min" in spec and "max" in spec else "") + (f" ({title})" if title else "")
    elif typ == "date":
        dv = DataValidation(type="date", operator="greaterThanOrEqual", formula1=str(DATE_MIN), allow_blank=True)
        dv.error = "Enter a date (dd-mm-yyyy)"
    else: return
    dv.showErrorMessage = True; dv.errorTitle = "Check this value"
    ws.add_data_validation(dv); dv.add(ref)


def as_cell_value(v, typ):
    if v is None: return None
    if typ == "date" and isinstance(v, str):
        try: return dt.datetime.strptime(v, "%d-%m-%Y")
        except ValueError: return v
    return v if not isinstance(v, (list, dict)) else str(v)


def draw_paper(ws, mapping, data=None, ids=None, wb=None):
    pg = mapping.get("page") or {}
    for col, w in (pg.get("widths") or {}).items(): ws.column_dimensions[col].width = w
    for r, h in (pg.get("heights") or {}).items(): ws.row_dimensions[int(r)].height = h
    ws.sheet_view.showGridLines = False
    if pg.get("orientation") == "landscape": ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    for it in mapping.get("layout") or []:
        if it.get("style") == "fingerprint":  # recognises the sheet on upload; kept out of sight in a hidden row
            c = ws[it["at"]]; c.value = it["text"]; c.font = Font(size=6, color="FFFFFF")
            ws.row_dimensions[c.row].hidden = True; continue
        place(ws, it["at"], it.get("to"), it.get("text"), it.get("style") or "label")
        if it.get("h"): ws.row_dimensions[ws[it["at"]].row].height = it["h"]
    tables = {}
    for f in mapping.get("fields") or []:
        if f.get("by") == "const": continue
        if f.get("type") == "table": tables.setdefault(f["at"], []).append(f); continue
        typ = f.get("type")
        vc = place(ws, f["cell"], f.get("to"), None, None, fill=INPUT, locked=False)
        if typ == "text": vc.number_format = "@"
        elif typ == "date": vc.number_format = "dd-mm-yyyy"
        validate_cells(ws, f"{f['cell']}:{f.get('to') or f['cell']}", f, typ)
        if data is not None or ids is not None:
            v = get(ids or {}, f["field"][5:]) if f["field"].startswith("@ids.") else get(data or {}, f["field"])
            if v is not None: vc.value = as_cell_value(v, typ)
        if f.get("name") and wb is not None:
            wb.defined_names[f["name"]] = DefinedName(f["name"], attr_text=f"'{ws.title}'!${get_column_letter(vc.column)}${vc.row}")
    for at, fs in tables.items(): (draw_tgrid if fs[0].get("orient") == "columns" else draw_table_paper)(ws, at, fs, data)
    ws.protection.sheet = True  # printed text, labels and formulas stay put; the input cells are unlocked
    for k in ("formatColumns", "formatRows", "formatCells"): setattr(ws.protection, k, False)


def stored_rows(f, val):
    """Stored value -> the sheet rows it came from (the inverse of read_table / table_result)."""
    if val is None: return []
    how = f.get("rows_as", "list"); n = len(f.get("take") or f["columns"])
    if "pick" in f: return []  # handled by the caller, column by column
    if how == "dict": rows = [[k] + (list(v) if n > 2 else [v]) for k, v in val.items()]
    elif how == "values": rows = [[x] for x in val]
    elif how == "row": rows = [[f.get("key")] + (list(val) if n > 2 else [val])]
    elif how == "first": rows = [list(val)]
    else: rows = [list(x) for x in val]
    if f.get("keys_as"):
        back = {v: k for k, v in f["keys_as"].items()}
        rows = [[back.get(x[0], x[0])] + x[1:] for x in rows]
    if not f.get("take"): return rows
    out = []
    for x in rows:
        sheet = [[None] * len(c["group"]) if "group" in c else None for c in f["columns"]]
        for k, t in enumerate(f["take"]):
            if k >= len(x): break
            if isinstance(t, list): sheet[t[0]][t[1]] = x[k]
            else: sheet[t] = x[k]
        out.append(sheet)
    return out


def merged_rows(fs, data, labels):
    """The rows of a table region filled from every field that reads it (by printed row label, else by position)."""
    f0 = fs[0]; rows = {}
    def slot(x, i):
        if not labels: return i
        lab = x[0] if x else None
        if lab is None: return i  # this field does not read the label column: rows follow each other from the top
        hit = next((n for n, l in enumerate(labels) if str(l).strip().lower() == str(lab).strip().lower()), None)
        return hit if hit is not None else len(labels) + i
    for f in fs:
        val = get(data, f["field"]) if data is not None else None
        if val is None: continue
        if "pick" in f:
            for i, x in enumerate(val): rows.setdefault(i, [[None] * len(c["group"]) if "group" in c else None for c in f0["columns"]])[f["pick"]] = x
            continue
        for i, x in enumerate(stored_rows(f, val)):
            cur = rows.setdefault(slot(x, i), [[None] * len(c["group"]) if "group" in c else None for c in f0["columns"]])
            for p, v in enumerate(x):
                if isinstance(v, list): cur[p] = [a if a is not None else (cur[p][q] if isinstance(cur[p], list) and q < len(cur[p]) else None) for q, a in enumerate(v)]
                elif v is not None: cur[p] = v
    return rows


def calc_formula(expr, ref):
    """'{Ambient 1}' in a column's formula -> that column's cell in the same row (or, sideways, the same record)."""
    return re.sub(r"\{([^}]+)\}", lambda m: ref(m.group(1)), expr)


def draw_table_paper(ws, at, fs, data):
    f0 = fs[0]; _, r0, c0 = addr(at); specs = col_specs(f0); labels = f0.get("row_labels")
    two = any(cs.get("title") for *_, cs in specs)
    if f0.get("caption"): place(ws, f"{get_column_letter(c0)}{r0 - 1}", f"{get_column_letter(c0 + len(specs) - 1)}{r0 - 1}", f0["caption"], "section")
    letters, c, done = {}, c0, set()
    for pos, j, lab, typ, cs in specs:
        letters[lab] = get_column_letter(c)
        hr = r0 + 1 if (two and j is not None) else r0
        place(ws, f"{get_column_letter(c)}{hr}", None, lab, "head")
        if two and j is None: ws.merge_cells(start_row=r0, start_column=c, end_row=r0 + 1, end_column=c)
        if two and j == 0 and pos not in done:
            n = len(f0["columns"][pos]["group"]); done.add(pos)
            place(ws, f"{get_column_letter(c)}{r0}", f"{get_column_letter(c + n - 1)}{r0}" if n > 1 else None, cs.get("title") or "", "head")
        if cs.get("w"): ws.column_dimensions[get_column_letter(c)].width = max(ws.column_dimensions[get_column_letter(c)].width or 0, cs["w"])
        c += 1
    first = r0 + (2 if two else 1)
    longest = max(len(str(lab)) for *_, lab, _t, _c in specs)  # wrapped headers get the height they need
    if longest > 10: ws.row_dimensions[first - 1].height = 15 * min(4, 1 + longest // 11)
    rows = merged_rows(fs, data, labels)
    n = max(len(labels) if labels else f0.get("rows", 5), (max(rows) + 1) if rows else 0)
    for i in range(n):
        r = first + i; out = rows.get(i); cc = c0
        for pos, j, lab, typ, cs in specs:
            cell = ws.cell(r, cc); cell.border = BOX
            v = None if out is None else out[pos] if j is None else (out[pos][j] if isinstance(out[pos], list) and j < len(out[pos]) else None)
            if pos == 0 and labels:
                cell.value = labels[i] if i < len(labels) else v; cell.fill = FIXED; cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif cs.get("calc"):
                cell.fill = CALC; cell.number_format = cs.get("format", "0.00")
                cell.value = v if v is not None else calc_formula(cs["calc"], lambda l: f"{letters[l]}{r}")
            else:
                cell.fill = INPUT; cell.protection = Protection(locked=False)
                if typ == "text": cell.number_format = "@"
                if v is not None: cell.value = as_cell_value(v, typ)
            cc += 1
    for k, (pos, j, lab, typ, cs) in enumerate(specs):  # one validation per column
        if (pos == 0 and labels) or cs.get("calc"): continue
        L = get_column_letter(c0 + k); validate_cells(ws, f"{L}{first}:{L}{first + n - 1}", dict(cs, title=lab), typ)


def draw_tgrid(ws, at, fs, data):
    f0 = fs[0]; _, r0, c0 = addr(at); cols = f0["columns"]; labels = f0.get("row_labels")
    n = f0.get("records") or len(labels or []) or 6
    rowof = {c["label"]: r0 + i for i, c in enumerate(cols)}
    for i, c in enumerate(cols): place(ws, f"{get_column_letter(c0)}{r0 + i}", None, c["label"], "label_b")
    rows = merged_rows(fs, data, labels)
    for k in range(n):
        cc = c0 + 1 + k; L = get_column_letter(cc); out = rows.get(k)
        for i, c in enumerate(cols):
            cell = ws.cell(r0 + i, cc); cell.border = BOX; v = out[i] if out else None
            if i == 0 and labels:
                cell.value = labels[k] if k < len(labels) else v; cell.fill = FIXED; cell.font = Font(bold=True); cell.alignment = Alignment(horizontal="center")
            elif c.get("calc"):
                cell.fill = CALC; cell.number_format = c.get("format", "0.00"); cell.value = v if v is not None else calc_formula(c["calc"], lambda l: f"{L}{rowof[l]}")
            else:
                cell.fill = INPUT; cell.protection = Protection(locked=False)
                if c.get("type") == "text": cell.number_format = "@"
                if v is not None: cell.value = as_cell_value(v, c.get("type", "number"))
    for i, c in enumerate(cols):
        if (i == 0 and labels) or c.get("calc"): continue
        validate_cells(ws, f"{get_column_letter(c0 + 1)}{r0 + i}:{get_column_letter(c0 + n)}{r0 + i}", c, c.get("type", f0.get("cell_type", "number")))


# ------------------------------------------------------------------ comparing template versions
def diff(a, b):
    """Field-level differences between two mappings (for the administrator's review before activating)."""
    fa = {f["field"]: f for f in a.get("fields") or []}; fb = {f["field"]: f for f in b.get("fields") or []}
    out = [dict(field=k, change="removed") for k in fa if k not in fb] + [dict(field=k, change="added") for k in fb if k not in fa]
    for k in fa.keys() & fb.keys():
        ch = {x: (fa[k].get(x), fb[k].get(x)) for x in set(fa[k]) | set(fb[k]) if fa[k].get(x) != fb[k].get(x)}
        if ch: out.append(dict(field=k, change="changed", what=ch))
    if a.get("fingerprint") != b.get("fingerprint"): out.append(dict(field="(fingerprint)", change="changed", what={"fingerprint": (a.get("fingerprint"), b.get("fingerprint"))}))
    return sorted(out, key=lambda x: x["field"])


def validate_mapping(m):
    """Problems in a mapping an administrator wrote, before it is saved."""
    err = []
    if not isinstance(m, dict): return ["The mapping must be a JSON object"]
    if not isinstance(m.get("fields"), list) or not m["fields"]: err.append("The mapping needs a list of fields")
    seen = set()
    for n, f in enumerate(m.get("fields") or [], 1):
        if not isinstance(f, dict) or not f.get("field"): err.append(f"Field {n}: needs a 'field' path"); continue
        try: tokens(f["field"].lstrip("@"))
        except TemplateError as e: err.append(f"Field {n}: {e}")
        if f["field"] in seen and "pick" not in f: err.append(f"Field {n}: '{f['field']}' is mapped twice")
        seen.add(f["field"])
        if f.get("type") == "table":
            if not f.get("columns"): err.append(f"Field {n} ({f['field']}): a table needs columns")
            for c in f.get("columns") or []:
                if not isinstance(c, dict) or not (c.get("label") or c.get("group")): err.append(f"Field {n} ({f['field']}): every column needs a label or a group of labels")
            if f.get("at"):
                try: addr(f["at"])
                except TemplateError as e: err.append(f"Field {n}: {e}")
        elif f.get("by") != "const":
            if not (f.get("cell") or f.get("label") or f.get("name")): err.append(f"Field {n} ({f['field']}): give a cell, a label or a defined name")
            if f.get("cell"):
                try: addr(f["cell"])
                except TemplateError as e: err.append(f"Field {n}: {e}")
            if f.get("type") not in (None, "any", "number", "text", "date", "bool"): err.append(f"Field {n}: unknown type '{f.get('type')}'")
    fp = m.get("fingerprint")
    if not isinstance(fp, dict) or not any(k in fp for k in ("contains", "equals", "label", "sheet")):
        err.append("The mapping needs a fingerprint (contains / equals / label / sheet) so uploads can be recognised")
    return err


def grid(raw, name="sample.xlsx", max_rows=60, max_cols=26):
    """Cells of each sheet for the administrator's click-to-bind view."""
    b = Book(raw, name); out = []
    for ws in b.sheets(hidden=True):
        sh = Sheet(b, ws)
        rows = []
        for r in range(1, min(ws.max_row, max_rows) + 1):
            row = []
            for c in range(1, min(ws.max_column, max_cols) + 1):
                v = ws.cell(r, c).value
                row.append(None if v is None else v.strftime("%d-%m-%Y") if isinstance(v, (dt.date, dt.datetime)) else v if isinstance(v, (int, float, str)) else str(v))
            rows.append(row)
        names = {n: coord for n, (t, coord) in b.names.items() if t == ws.title}
        out.append(dict(name=ws.title, hidden=ws.sheet_state != "visible", rows=rows, merged=[str(m) for m in ws.merged_cells.ranges], names=names,
                        formulas=[sh.ref(c.row, c.column).split("!")[1] for row in b.f[ws.title].iter_rows(max_row=max_rows, max_col=max_cols) for c in row
                                  if isinstance(c.value, str) and c.value.startswith("=")]))
    return out
