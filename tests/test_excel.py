"""Excel extraction with templates."""
import copy, io, json, os, re, sqlite3, sys, unittest, zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from test_app import Base, DEMO, REQUEST, aletheia, signed_in, up  # noqa: E402
import paper_templates as P, seed_templates as S, xltemplates as X  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

LOGS = [k for k in S.SPECS]


def filled(sections=None, data=DEMO, ids=None):
    """A workbook of logsheets in the current layout, filled with the demo values."""
    ids = DEMO["ids"] if ids is None else ids
    return X.workbook([(S.layout(k), data[k], ids) for k in (sections or LOGS)])


def filled2(sections=None, data=DEMO, ids=None):
    """The same, in the paper layout (version 2, the active one)."""
    ids = DEMO["ids"] if ids is None else ids
    return X.workbook([(P.paper(k), data[k], ids) for k in (sections or LOGS)])


def present(d): return {k: v for k, v in d.items() if v is not None}


def with_cached(raw, cell, value):
    """Give a formula cell a saved value, as Excel does when it saves a workbook (openpyxl cannot)."""
    zin = zipfile.ZipFile(io.BytesIO(raw)); out = io.BytesIO(); zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        b = zin.read(it.filename)
        if it.filename.startswith("xl/worksheets/sheet"):
            b = re.sub(rf'(<c r="{cell}"[^>]*>\s*<f>[^<]*</f>)\s*<v\s*/>'.encode(), rb"\1<v>" + str(value).encode() + rb"</v>", b)
        zout.writestr(it, b)
    zout.close(); return out.getvalue()


class RoundTrip(Base):
    def test_every_logsheet_reads_back_exactly(self):
        for k in LOGS:
            r = X.extract(X.Book(filled([k])), S.layout(k))
            self.assertEqual(present(r["data"]), DEMO[k], k); self.assertEqual(r["errors"], [], k); self.assertEqual(r["warnings"], [], k)

    def test_one_workbook_with_every_test_fills_the_job(self):
        i = self.job(); raw = filled()
        r = self.c.post(f"/api/jobs/{i}/import", json=up("all_logsheets.xlsx", raw)); self.assertEqual(r.status_code, 200, r.json)
        d = self.get(i)["data"]
        for k in LOGS: self.assertEqual(present(d[k]), DEMO[k], k)
        self.assertEqual(d["ids"]["sc"], DEMO["ids"]["sc"])
        self.assertIn("read with template temp-std v1", " ".join(r.json["notes"]))
        m = self.get(i)["meta"]["temp"]; self.assertIsNotNone(m["template_id"])
        self.c.post(f"/api/jobs/{i}/validate"); self.assertEqual(self.gen(i).status_code, 200)

    def test_blank_sheets_and_job_export_round_trip(self):
        b = self.c.get("/api/logsheets/all.xlsx"); self.assertEqual(b.status_code, 200)
        wb = load_workbook(io.BytesIO(b.data)); self.assertEqual(len(wb.sheetnames), len(LOGS))
        self.assertIn("TR2_RHV_COLD", wb.defined_names)  # the paper layout (version 2) is the one handed out
        self.assertEqual(self.c.get("/api/logsheets/temp.xlsx").status_code, 200)
        i = self.c.post("/api/demo").json["id"]
        x = self.c.get(f"/api/jobs/{i}/logsheets.xlsx"); self.assertEqual(x.status_code, 200)
        j = self.job("CPRIBLRSCL25T1700"); self.assertEqual(self.c.post(f"/api/jobs/{j}/import", json=up("export.xlsx", x.data)).status_code, 200)
        a, b2 = self.get(i)["data"], self.get(j)["data"]
        for k in LOGS: self.assertEqual(present(b2[k]), a[k], k)


class Preview(Base):
    def test_preview_shows_cells_and_stores_nothing(self):
        i = self.job(); p = self.c.post(f"/api/jobs/{i}/excel/preview", json=up("nl.xlsx", filled(["noload", "sc"])))
        self.assertEqual(p.status_code, 200, p.json); self.assertTrue(p.json["ok"])
        nl = next(s for s in p.json["sheets"] if s["section"] == "noload")
        v = next(f for f in nl["fields"] if f["field"] == "v100")
        self.assertEqual((v["value"], v["status"], v["by"]), (433.06, "ok", "name")); self.assertRegex(v["cell"], r"!B\d+$")
        rows = next(f for f in nl["fields"] if f["field"] == "rows"); self.assertIn("3 rows", rows["note"])
        self.assertFalse(nl["series_matches"])  # the demo no-load sheet has 25T1656 written on it
        self.assertTrue(next(s for s in p.json["sheets"] if s["section"] == "sc")["series_matches"])
        self.assertNotIn("noload", self.get(i)["data"])
        r = self.c.post(f"/api/jobs/{i}/excel/import", json=up("nl.xlsx", filled(["noload", "sc"])))
        self.assertEqual(r.status_code, 200); self.assertEqual(present(self.get(i)["data"]["noload"]), DEMO["noload"])

    def test_required_value_missing_blocks_the_section(self):
        d = copy.deepcopy(DEMO); d["proforma"]["kva"] = None; d["proforma"]["z_pct"] = None
        i = self.job(); p = self.c.post(f"/api/jobs/{i}/excel/preview", json=up("pf.xlsx", filled(["proforma"], data=d))).json
        self.assertFalse(p["ok"]); errs = " ".join(p["sheets"][0]["errors"]); self.assertIn("Rating (kVA)", errs); self.assertIn("required", errs)
        r = self.c.post(f"/api/jobs/{i}/import", json=up("pf.xlsx", filled(["proforma"], data=d)))
        self.assertEqual(r.status_code, 400); self.assertNotIn("proforma", self.get(i)["data"])  # never stored silently as NA

    def test_series_on_sheet_that_does_not_match_the_job_is_flagged(self):
        i = self.job("CPRIBLRSCL25T1999"); p = self.c.post(f"/api/jobs/{i}/excel/preview", json=up("x.xlsx", filled(["sc"]))).json
        self.assertFalse(p["sheets"][0]["series_matches"])


class Drift(Base):
    def shifted(self, k, rows=2, cols=0, names=False, relabel=None):
        """The same sheet as the lab might keep it: everything moved, no defined names, a label reworded."""
        m = copy.deepcopy(S.layout(k))
        for f in m["fields"]:
            for key in ("cell", "at"):
                if f.get(key):
                    _, r, c = X.addr(f[key]); f[key] = f"{X.get_column_letter(c + cols)}{r + rows}"
            if not names: f.pop("name", None)
            if relabel and f.get("label") in relabel: f["label"] = relabel[f["label"]]
        return X.workbook([(m, DEMO[k], DEMO["ids"])])

    def test_inserted_rows_and_columns_without_names_are_found_by_label(self):
        for k in ("noload", "temp", "routine", "resistance", "pressure"):
            r = X.extract(X.Book(self.shifted(k, rows=3, cols=1)), S.layout(k))
            self.assertEqual(present(r["data"]), DEMO[k], k)
            self.assertTrue(any(f.get("by") == "label" for f in r["fields"]), k)

    def test_reworded_label_is_matched_fuzzily(self):
        raw = self.shifted("noload", relabel={"Rated voltage, 100% (V)": "Rated voltages, 100% (V)"})
        r = X.extract(X.Book(raw), S.layout("noload")); self.assertEqual(r["data"]["v100"], 433.06)

    def test_duplicate_label_is_reported_not_guessed(self):
        raw = self.shifted("noload", rows=0)
        wb = load_workbook(io.BytesIO(raw)); ws = wb.active; ws["H30"] = "Remarks"; ws["I30"] = "copied by mistake"; ws["H31"] = "Remarks"
        out = io.BytesIO(); wb.save(out)
        m = S.layout("noload")
        for f in m["fields"]:  # a mapping written for the lab's own sheet: by label only, no cell to prefer
            if f["field"] == "remarks": f["by"] = ["label"]; f.pop("cell")
        r = X.extract(X.Book(out.getvalue()), m)
        f = next(x for x in r["fields"] if x["field"] == "remarks"); self.assertEqual(f["status"], "ambiguous")
        self.assertTrue(r["errors"])

    def test_formulas_are_taken_as_logged_or_refused_when_never_calculated(self):
        wb = load_workbook(io.BytesIO(filled(["noload"]))); ws = wb.active
        cell = wb.defined_names["NL_V100"].attr_text.split("!")[1].replace("$", "")
        ws[cell] = "=433+0.06"; out = io.BytesIO(); wb.save(out)
        r = X.extract(X.Book(out.getvalue()), S.layout("noload"))
        f = next(x for x in r["fields"] if x["field"] == "v100"); self.assertEqual(f["status"], "formula_unsaved"); self.assertTrue(r["errors"])
        r = X.extract(X.Book(with_cached(out.getvalue(), cell, 433.06)), S.layout("noload"))
        f = next(x for x in r["fields"] if x["field"] == "v100"); self.assertEqual((f["status"], f["value"]), ("formula", 433.06)); self.assertFalse(r["errors"])

    def test_decimal_comma_only_by_explicit_rule(self):
        wb = load_workbook(io.BytesIO(filled(["noload"]))); ws = wb.active
        cell = wb.defined_names["NL_V100"].attr_text.split("!")[1].replace("$", ""); ws[cell] = "433,06"; out = io.BytesIO(); wb.save(out)
        r = X.extract(X.Book(out.getvalue()), S.layout("noload"))
        f = next(x for x in r["fields"] if x["field"] == "v100"); self.assertEqual(f["status"], "type"); self.assertIn("decimal comma", f["note"])
        m = S.layout("noload"); next(x for x in m["fields"] if x["field"] == "v100")["decimal"] = ","
        self.assertEqual(X.extract(X.Book(out.getvalue()), m)["data"]["v100"], 433.06)

    def test_dates_hidden_sheets_and_refused_files(self):
        import datetime as dt
        wb = load_workbook(io.BytesIO(filled(["sc"]))); ws = wb.active
        cell = wb.defined_names["SC_DATE"].attr_text.split("!")[1].replace("$", ""); ws[cell] = dt.datetime(2025, 10, 31)
        h = wb.create_sheet("Old copy"); h.sheet_state = "hidden"; h["A1"] = "Aletheia template noload"
        out = io.BytesIO(); wb.save(out); b = X.Book(out.getvalue())
        self.assertEqual(X.extract(b, S.layout("sc"))["data"]["date"], "31-10-2025")  # a real Excel date
        self.assertEqual([t["section"] for _, t in X.detect(b, [dict(mapping=S.layout(k), section=k) for k in LOGS])], ["sc"])  # hidden sheet ignored
        i = self.job()
        for name, data, msg in (("x.xls", b"\xd0\xcf\x11\xe0" + b"0" * 100, "save it as .xlsx"), ("x.xlsm", filled(["sc"]), "Macro-enabled"),
                                ("x.xlsx", b"PK\x03\x04 broken", "damaged")):
            r = self.c.post(f"/api/jobs/{i}/excel/preview", json=up(name, data)); self.assertEqual(r.status_code, 400, name); self.assertIn(msg, r.json["error"][0])


class ManyReports(Base):
    def test_records_workbook_covers_many_reports(self):
        for s in ("CPRIBLRSCL25T1654", "CPRIBLRSCL25T1655", "CPRIBLRSCL25T1656"):
            i = self.job(s); self.c.post(f"/api/jobs/{i}/import", json=up("all.xlsx", filled(["proforma", "temp"])))
        x = self.c.get("/api/records.xlsx"); self.assertEqual(x.status_code, 200)
        wb = load_workbook(io.BytesIO(x.data)); self.assertEqual(wb.sheetnames, ["Jobs", "Tests", "Key values (as logged)"])
        self.assertEqual(wb["Jobs"].max_row, 4)
        vals = wb["Key values (as logged)"]; self.assertEqual(vals["B2"].value, 250); self.assertEqual(vals["F2"].value, 26.01)
        one = load_workbook(io.BytesIO(self.c.get(f"/api/records.xlsx?ids={i}").data)); self.assertEqual(one["Jobs"].max_row, 2)
        self.assertEqual(self.c.get("/api/records.xlsx?q=1655").status_code, 200)


class Registry(Base):
    def tid(self, key):
        return next(t["id"] for t in self.c.get(f"/api/templates?key={key}&status=active").json)

    def test_new_version_tested_activated_and_old_jobs_unchanged(self):
        i = self.job(); self.c.post(f"/api/jobs/{i}/import", json=up("nl.xlsx", filled(["noload"])))
        old_tid = self.get(i)["meta"]["noload"]["template_id"]; old = self.c.get(f"/api/templates/{old_tid}").json
        m = copy.deepcopy(old["mapping"])
        next(f for f in m["fields"] if f["field"] == "v100")["label"] = "Applied voltage, 100% (V)"
        self.assertEqual(self.c.post("/api/templates", json=dict(from_id=old_tid, mapping=m)).status_code, 403)  # testers do not manage templates
        r = self.admin.post("/api/templates", json=dict(from_id=old_tid, mapping=m)); self.assertEqual(r.status_code, 201, r.json)
        new = r.json["id"]; self.assertEqual(r.json["version"], 3)  # versions 1 (read this file) and 2 (the paper layout) exist
        self.assertEqual(self.admin.post("/api/templates", json=dict(from_id=old_tid, mapping={"fields": []})).status_code, 400)
        d = self.admin.get(f"/api/templates/{old_tid}/diff/{new}").json; self.assertEqual([x["field"] for x in d], ["v100"])
        past = self.admin.post(f"/api/templates/{new}/test-past", json={}).json
        self.assertEqual(past["files"][0]["result"], "same as stored")  # names still find the cells: the change is safe
        live = self.admin.post("/api/templates/test", json=dict(up("nl.xlsx", filled(["noload"])), mapping=m)).json
        self.assertEqual(live["data"]["v100"], 433.06)
        self.assertEqual(self.admin.post(f"/api/templates/{new}/activate").status_code, 200)
        st = {t["version"]: t["status"] for t in self.admin.get("/api/templates?key=noload-std").json}; self.assertEqual(st, {1: "retired", 2: "retired", 3: "active"})
        with self.assertRaises(sqlite3.IntegrityError):  # a template that has read data is frozen
            with aletheia.db() as c: c.execute("UPDATE templates SET mapping='{}' WHERE id=?", (old_tid,))
        self.assertEqual(self.get(i)["meta"]["noload"]["template_id"], old_tid)  # the old job keeps the version that read it
        x = load_workbook(io.BytesIO(self.c.get(f"/api/jobs/{i}/logsheets.xlsx").data)); self.assertIn("Rated voltage, 100% (V)", [c.value for c in x.active["A"]] + [c.value for c in x.active["D"]])
        j = self.job("CPRIBLRSCL25T1700"); self.c.post(f"/api/jobs/{j}/import", json=up("nl.xlsx", filled(["noload"])))
        self.assertEqual(self.get(j)["meta"]["noload"]["template_id"], new)
        e = self.admin.get(f"/api/templates/{new}/export"); self.assertEqual(json.loads(e.data)["version"], 3)
        self.assertEqual(self.admin.post("/api/templates/import", json=json.loads(e.data)).json["version"], 4)

    def test_admin_tools_grid_and_sample(self):
        tid = self.tid("temp-std")
        g = self.admin.post("/api/templates/grid", json=up("t.xlsx", filled2(["temp"]))).json["sheets"][0]
        self.assertEqual(g["rows"][0][0], "Aletheia template temp v2"); self.assertIn("tr2_rhv_cold", g["names"])
        d = self.admin.post("/api/templates", json=dict(from_id=tid)).json["id"]
        self.assertEqual(self.admin.post(f"/api/templates/{d}/sample", json=up("t.xlsx", filled2(["temp"]))).status_code, 200)
        r = self.admin.post("/api/templates/test", json=dict(template_id=d, mapping=self.admin.get(f"/api/templates/{d}").json["mapping"])).json
        self.assertEqual({k: r["data"][k] for k in DEMO["temp"]}, DEMO["temp"]); self.assertTrue(r["fingerprint_found"])
        self.assertEqual(self.admin.post(f"/api/templates/{d}/retire").status_code, 200)  # a draft is simply discarded
        self.assertEqual(self.admin.get(f"/api/templates/{d}").status_code, 404)


class RequestForm(Base):
    def test_request_form_is_filled_in_online_only(self):
        # the Excel request form stays in the template registry, but a request is raised by the customer online: staff cannot
        # read one in, and customers cannot upload one instead of filling in the form
        self.assertEqual(self.c.get("/api/request-form.xlsx").status_code, 200)
        self.assertEqual(self.c.post("/api/intake/from-excel", json=up("filled form.xlsx", b"PK")).status_code, 404)
        self.assertIn(self.cust.post("/api/customer/request-forms", json=up("filled form.xlsx", b"PK")).status_code, (404, 405))

    def fill(self, raw, values):
        """The downloaded form filled in as a customer would, cell by named cell."""
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw))
        for name, v in values.items():
            sheet, ref = next(iter(wb.defined_names[name].destinations))
            wb[sheet][ref.replace("$", "")].value = v
        out = io.BytesIO(); wb.save(out); return out.getvalue()

    def test_excel_form_fills_the_online_form(self):
        # the Excel sheet laid out as the paper form fills the online form; the customer still checks it and sends it online
        r = self.cust.get("/api/request-form.xlsx"); self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.data)).worksheets[0]
        printed = " ".join(str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str))
        for text in ("CENTRAL POWER RESEARCH INSTITUTE", "Sheet 1 of 3", "Sheet 2 of 3", "Sheet 3 of 3", "Rating of the sample(s) to be tested", "MSME Discount"):
            self.assertIn(text, printed)
        cells = {f"RQ_{k.upper()}": ("Yes" if v is True else v) for k, v in REQUEST.items() if v != ""}
        cells.update(RQ_DRAWINGS="Not applicable: no drawings for a standard design", RQ_PLAN_SC="Yes", RQ_PLAN_TEMP="Yes", RQ_PLAN_PRESSURE="No")
        got = self.cust.post("/api/customer/requests/excel", json=up("my request.xlsx", self.fill(r.data, cells)))
        self.assertEqual(got.status_code, 200, got.json); got = got.json
        self.assertEqual(got["plan"], ["sc", "temp"]); self.assertEqual(got["na"], {"drawings": "no drawings for a standard design"})
        self.assertEqual(got["values"]["pin"], "600058"); self.assertEqual(got["values"]["samples"], "1")
        self.assertTrue(got["values"]["declare_terms"]); self.assertTrue(got["values"]["decision_rule"].startswith("(i) Decision on compliance"))
        self.assertEqual(got["problems"], [])
        chk = self.cust.post("/api/customer/requests/check", json=dict(got["values"], na=got["na"], plan=got["plan"])).json
        self.assertTrue(chk["ok"], chk)  # what the sheet gave passes the online form's own checks
        with aletheia.db() as c: self.assertEqual(c.execute("SELECT COUNT(*) FROM customer_forms").fetchone()[0], 0)  # nothing is sent by reading

    def test_excel_form_only_for_the_request_form_and_customers(self):
        r = self.cust.get("/api/request-form.xlsx").data
        self.assertEqual(self.c.post("/api/customer/requests/excel", json=up("f.xlsx", r)).status_code, 403)
        bad = self.cust.post("/api/customer/requests/excel", json=up("logsheet.xlsx", self.c.get("/api/logsheets/sc.xlsx").data))
        self.assertEqual(bad.status_code, 400); self.assertIn("not the Customer Request Form", bad.json["error"][0])


if __name__ == "__main__":
    unittest.main()
