"""Excel templates laid out as the paper logsheets (version 2, paper_templates.py; item 7 of the feedback of 10-11 Oct 2026)."""
import io, json, os, sqlite3, sys, unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from test_app import Base, DEMO, aletheia, up  # noqa: E402
from test_excel import filled, with_cached  # noqa: E402  (version 1 workbooks, and Excel's saved formula values)
import excel_routes, paper_templates as P, xltemplates as X  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

LOGS = list(P.BUILD)


def paper(sections=None, data=DEMO, ids=None):
    """A workbook of paper-layout logsheets filled with the demo values."""
    return X.workbook([(P.paper(k), data[k], DEMO["ids"] if ids is None else ids) for k in (sections or LOGS)])


def same(stored, read):
    """Every stored value of the demo comes back; fields the paper adds and nothing was written in are left empty."""
    strip = lambda v: {a: strip(b) for a, b in v.items() if b is not None} if isinstance(v, dict) else v
    return {k: strip(read.get(k)) for k in stored} == stored


class PaperLayout(Base):
    def test_every_paper_logsheet_reads_back_exactly(self):
        for k in LOGS:
            r = X.extract(X.Book(paper([k])), P.paper(k))
            self.assertTrue(same(DEMO[k], r["data"]), k); self.assertEqual((r["errors"], r["warnings"]), ([], []), k)
            if k in DEMO["ids"] and k != "work": self.assertEqual(r["ids"][k], DEMO["ids"][k], k)

    def test_the_paper_layout_is_handed_out_and_version_1_still_imports(self):
        for k in LOGS:
            st = {t["version"]: (t["status"], t["note"]) for t in self.admin.get(f"/api/templates?key={k}-std").json}
            self.assertEqual(st[1][0], "retired", k); self.assertEqual(st[2], ("active", excel_routes.PAPER_NOTE), k)
        i = self.job(); r = self.c.post(f"/api/jobs/{i}/import", json=up("paper.xlsx", paper()))
        self.assertEqual(r.status_code, 200, r.json); self.assertIn("read with template temp-std v2", " ".join(r.json["notes"]))
        d = self.get(i)["data"]
        for k in LOGS: self.assertTrue(same(DEMO[k], d[k]), k)
        self.c.post(f"/api/jobs/{i}/validate"); self.assertEqual(self.gen(i).status_code, 200)  # checks and report work on it
        j = self.job("CPRIBLRSCL25T1700"); r = self.c.post(f"/api/jobs/{j}/import", json=up("old.xlsx", filled(["temp", "sc"])))
        self.assertEqual(r.status_code, 200, r.json); self.assertIn("read with template temp-std v1", " ".join(r.json["notes"]))  # filled in earlier

    def test_blank_sheet_is_protected_validated_and_carries_a_hidden_fingerprint(self):
        wb = load_workbook(io.BytesIO(self.c.get("/api/logsheets/sc.xlsx").data)); ws = wb.active
        self.assertEqual(ws.title, "Short-Circuit Withstand Test")
        self.assertEqual(ws["A1"].value, "Aletheia template sc v2"); self.assertTrue(ws.row_dimensions[1].hidden)
        self.assertTrue(ws.protection.sheet)
        m = P.paper("sc"); f = next(x for x in m["fields"] if x["field"] == "date")
        self.assertFalse(ws[f["cell"]].protection.locked)  # an input cell
        title = next(x for x in m["layout"] if x["style"] == "title"); self.assertTrue(ws[title["at"]].protection.locked)  # printed text
        kinds = {(d.type, d.formula1) for d in ws.data_validations.dataValidation}
        self.assertIn(("list", '"NT,HT,LT"'), kinds); self.assertIn("date", {t for t, _ in kinds}); self.assertIn("decimal", {t for t, _ in kinds})
        shots = next(x for x in m["fields"] if x["field"] == "shots"); _, r0, c0 = X.addr(shots["at"])
        avg = ws.cell(r0 + 2, c0 + next(i for i, c in enumerate(X.col_specs(shots)) if c[2] == "RMS (Avg)"))
        self.assertTrue(str(avg.value).startswith("=IF(COUNT(")); self.assertTrue(avg.protection.locked)  # calculated by the sheet

    def test_calculated_cells_need_the_value_excel_saved(self):
        m = P.paper("noload"); f = next(x for x in m["fields"] if x["field"] == "rows"); _, r0, c0 = X.addr(f["at"])
        wb = load_workbook(io.BytesIO(X.workbook([(m, None, None)]))); ws = wb.active; r = r0 + 2  # the BT row, under two header rows
        for k, v in enumerate([251.133, 252.814, 246.166, None, 1.907, 1.445, 1.714, None, 109.18, 84.52, 189.22, None, 49.95, 383.18]):
            if v is not None: ws.cell(r, c0 + 1 + k).value = v
        out = io.BytesIO(); wb.save(out); raw = out.getvalue()
        res = X.extract(X.Book(raw), m); self.assertIn("formula without a saved value", " ".join(res["errors"]))  # never calculated: refused
        for col, v in ((4, 250.037667), (8, 1.688667), (12, 382.92)):
            raw = with_cached(raw, f"{X.get_column_letter(c0 + col)}{r}", v)
        res = X.extract(X.Book(raw), m); self.assertEqual(res["errors"], [])
        self.assertEqual(res["data"]["rows"], [["BT", [251.133, 252.814, 246.166], 250.037667, [1.907, 1.445, 1.714], 1.688667, [109.18, 84.52, 189.22], 382.92, 49.95, 383.18]])

    def test_required_readings_missing_block_and_blank_label_rows_are_skipped(self):
        d = json.loads(json.dumps(DEMO)); d["sc"]["shots"] = None
        r = X.extract(X.Book(paper(["sc"], data=d)), P.paper("sc")); self.assertIn("S.C. Test results: table is empty (required)", " ".join(r["errors"]))
        r = X.extract(X.Book(paper(["noload"])), P.paper("noload"))
        self.assertEqual([x[0] for x in r["data"]["rows"]], ["BT", "AT", "112.5%"])  # 90 % and 110 % printed, nothing logged

    def test_tester_downloads_a_blank_sheet_with_the_jobs_numbers_on_it(self):
        i = self.job(); j = self.get(i)
        r = self.c.get(f"/api/jobs/{i}/logsheets/sc.xlsx"); self.assertEqual(r.status_code, 200)  # self.c is a tester
        self.assertIn(f"{j['series']}_sc_logsheet_v2", r.headers["Content-Disposition"])
        res = X.extract(X.Book(r.data), P.paper("sc"))
        self.assertEqual(res["ids"]["sc"], [j["series"], j["sample"]])
        self.assertIsNone(res["data"].get("shots")); self.assertIn("table is empty (required)", " ".join(res["errors"]))  # readings left to the tester
        w = X.extract(X.Book(self.c.get(f"/api/jobs/{i}/logsheets/work.xlsx").data), P.paper("work"))["data"]
        self.assertEqual((w["series"], w["customer"]), (j["series"], j["customer"]))
        self.assertEqual(self.c.get(f"/api/jobs/{i}/logsheets/request.xlsx").status_code, 404)
        self.assertEqual(self.cust.get(f"/api/jobs/{i}/logsheets/sc.xlsx").status_code, 403)

    def test_preview_names_each_value_as_printed(self):
        i = self.job(); p = self.c.post(f"/api/jobs/{i}/excel/preview", json=dict(up("sc.xlsx", paper(["sc"])), section="sc")).json
        self.assertTrue(p["ok"], p); f = {x["field"]: x for x in p["sheets"][0]["fields"]}
        self.assertEqual(f["date"]["label"], "Date of Test"); self.assertEqual(f["shots"]["label"], "S.C. Test results")
        self.assertRegex(f["date"]["cell"], r"!J\d+$")

    def test_an_administrators_own_version_is_kept_and_the_paper_layout_offered_as_a_draft(self):
        c = sqlite3.connect(":memory:"); c.row_factory = sqlite3.Row
        c.executescript("CREATE TABLE sections(id INT); CREATE TABLE section_history(id INT);")
        excel_routes.init_db(c)  # a fresh database: version 1 seeded, then the paper layout made active
        self.assertEqual(c.execute("SELECT status FROM templates WHERE key='temp-std' AND version=2").fetchone()[0], "active")
        c2 = sqlite3.connect(":memory:"); c2.row_factory = sqlite3.Row
        c2.executescript("CREATE TABLE sections(id INT); CREATE TABLE section_history(id INT);")
        excel_routes.init_db.__globals__["paper_rollout"], keep = (lambda c: None), excel_routes.paper_rollout
        try: excel_routes.init_db(c2)  # an older installation: version 1 only
        finally: excel_routes.init_db.__globals__["paper_rollout"] = keep
        c2.execute("UPDATE templates SET status='retired' WHERE key='temp-std'")
        c2.execute("INSERT INTO templates(key,kind,section,name,version,status,mapping,note) VALUES('temp-std','logsheet','temp','Lab sheet',2,'active','{}','the lab''s own')")
        excel_routes.paper_rollout(c2)
        st = {r[0]: r[1] for r in c2.execute("SELECT version, status FROM templates WHERE key='temp-std'")}
        self.assertEqual(st, {1: "retired", 2: "active", 3: "draft"})  # the lab's version stays in use
        self.assertEqual(c2.execute("SELECT status FROM templates WHERE key='sc-std' AND version=2").fetchone()[0], "active")
        excel_routes.paper_rollout(c2)  # once only
        self.assertEqual(c2.execute("SELECT COUNT(*) FROM templates WHERE key='temp-std'").fetchone()[0], 3)


if __name__ == "__main__":
    unittest.main()
