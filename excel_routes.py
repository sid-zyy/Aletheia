"""Template registry and Excel routes (docs/NEXT_STEPS.md section 5). Installed into the app by install(app_module).

* Registry: versioned templates (draft -> active -> retired), seeded once from seed_templates.py, then managed by Admin.
* Upload: a workbook is matched sheet by sheet to the active templates; the preview shows every value with its source cell;
  the import stores exactly what was previewed, with the template version that read it.
* Multi-report workbooks: one workbook holding several tests, or the sheets of several jobs (routed by the series number
  written on each sheet).
* Downloads: blank logsheets (one test or all in one workbook), the customer request form, a job's data as filled logsheets,
  and a register of many jobs in one workbook.
"""
import datetime as dt, io, json, re
from flask import jsonify, request, send_file, abort
import auth, integrity, xltemplates as X, seed_templates

A = None  # the running app module (app.py), set by install()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def init_db(c):
    c.executescript("""
    CREATE TABLE IF NOT EXISTS templates(id INTEGER PRIMARY KEY, key TEXT NOT NULL, kind TEXT NOT NULL, section TEXT, name TEXT, version INT NOT NULL,
        status TEXT NOT NULL DEFAULT 'draft', mapping TEXT NOT NULL, sample_name TEXT, sample BLOB, created_by INT, created_at TEXT,
        activated_by INT, activated_at TEXT, retired_at TEXT, note TEXT, UNIQUE(key, version));""")
    have = {r[1] for r in c.execute("PRAGMA table_info(sections)")}
    if "template_id" not in have: c.execute("ALTER TABLE sections ADD COLUMN template_id INT")
    have = {r[1] for r in c.execute("PRAGMA table_info(section_history)")}
    if "template_id" not in have: c.execute("ALTER TABLE section_history ADD COLUMN template_id INT")
    if not c.execute("SELECT 1 FROM templates").fetchone():
        t = dt.datetime.now().isoformat(timespec="seconds")
        for key, kind, section, name, m in seed_templates.all_templates():
            c.execute("INSERT INTO templates(key,kind,section,name,version,status,mapping,created_at,activated_at,note) VALUES(?,?,?,?,1,'active',?,?,?,?)",
                      (key, kind, section, name, json.dumps(m), t, t, "Version 1 shipped with Aletheia"))
    # a template that has read stored data is part of the record: never changed afterwards (a new version is made instead)
    c.execute("DROP TRIGGER IF EXISTS templates_frozen")
    c.execute("CREATE TRIGGER templates_frozen BEFORE UPDATE OF mapping ON templates WHEN OLD.status!='draft' "
              "BEGIN SELECT RAISE(ABORT, 'only a draft template can be changed; create a new version'); END")
    c.execute("DROP TRIGGER IF EXISTS templates_no_delete")
    c.execute("CREATE TRIGGER templates_no_delete BEFORE DELETE ON templates WHEN OLD.status!='draft' "
              "BEGIN SELECT RAISE(ABORT, 'templates are retired, never deleted'); END")


def row(r, full=False):
    d = {k: r[k] for k in r.keys() if k not in ("mapping", "sample")}
    m = json.loads(r["mapping"]); d["fields"] = len(m.get("fields") or []); d["has_sample"] = r["sample"] is not None
    if full: d["mapping"] = m
    return d


def active(c, kind=None, section=None):
    q, a = "SELECT * FROM templates WHERE status='active'", []
    if kind: q += " AND kind=?"; a.append(kind)
    if section: q += " AND section=?"; a.append(section)
    return [dict(id=r["id"], key=r["key"], version=r["version"], name=r["name"], section=r["section"], mapping=json.loads(r["mapping"]))
            for r in c.execute(q + " ORDER BY key", a)]


def template(c, tid):
    r = c.execute("SELECT * FROM templates WHERE id=?", (tid,)).fetchone()
    if not r: abort(404)
    return r


def file_arg(b, key="file"):
    f = b.get(key) or b
    name, raw = A.upload(f)
    if len(raw) > A.importers.MAX_BYTES: raise A.importers.ImportError_("File is too large (20 MB maximum)")
    return name, raw


def book(name, raw):
    try: return X.Book(raw, name)
    except X.TemplateError as e: raise A.importers.ImportError_(str(e)) from None


def norm_series(s): return re.sub(r"[^A-Z0-9]", "", str(s or "").upper()).replace("H", "4")


def read_workbook(c, name, raw, template_id=None):
    """Every recognised sheet of the workbook, extracted. Returns (sheets, unmatched sheet names)."""
    b = book(name, raw)
    ts = [dict(id=r["id"], key=r["key"], version=r["version"], name=r["name"], section=r["section"], mapping=json.loads(r["mapping"]))
          for r in [template(c, template_id)]] if template_id else active(c, "logsheet")
    found = X.detect(b, ts)
    if template_id and not found and b.sheets(): found = [(b.sheets()[0], ts[0])]  # the user chose the template explicitly
    out = []
    for ws, t in found:
        try: r = X.extract(b, t["mapping"], ws)
        except X.TemplateError as e: raise A.importers.ImportError_(f"Template {t['key']} v{t['version']}: {e}") from None
        r["template"] = {k: t[k] for k in ("id", "key", "version", "name")}
        sec = r["section"]
        r["series_on_sheet"] = (r["ids"].get(sec) or [None])[0] if r["ids"] else (r["data"].get("series") if sec == "work" else None)
        out.append(r)
    seen = {r["sheet"] for r in out}
    return out, [ws.title for ws in b.sheets() if ws.title not in seen]


def to_content(sheets):
    """Sections and identifiers from the extracted sheets of one job. Two sheets for the same test is an error, not a merge."""
    content, ids, used, dup = {}, {}, {}, []
    for r in sheets:
        s = r["section"]
        if s in content: dup.append(f"{A.NAMES.get(s, s)} appears on sheets '{used[s]}' and '{r['sheet']}'")
        content[s] = r["data"]; used[s] = r["sheet"]
        ids.update(r["ids"])
    if ids: content["ids"] = ids
    return content, dup


def create_template(c, b):
    """Insert a draft (inside the caller's write transaction). Returns {id, version} or a list of problems."""
    if b.get("from_id"):
        base = template(c, b["from_id"]); key, kind, section = base["key"], base["kind"], base["section"]
        m = b.get("mapping") or json.loads(base["mapping"]); name = b.get("name") or base["name"]
        sample = (base["sample_name"], base["sample"])
    else:
        key = re.sub(r"[^a-z0-9-]+", "-", str(b.get("key") or "").lower()).strip("-")
        kind, section, name, m, sample = b.get("kind") or "logsheet", b.get("section"), str(b.get("name") or "").strip(), b.get("mapping"), (None, None)
        if not key: return ["Give the template a key (e.g. noload-bay3)"]
        if kind not in ("logsheet", "request_form"): return ["Kind must be logsheet or request_form"]
        if kind == "logsheet" and section not in A.NAMES: return ["Choose which test document the template reads"]
        if kind == "request_form": section = "request"
        if c.execute("SELECT 1 FROM templates WHERE key=?", (key,)).fetchone(): return ["That key exists; create a new version of it instead"]
    err = X.validate_mapping(m)
    if err: return err
    v = c.execute("SELECT COALESCE(MAX(version),0)+1 FROM templates WHERE key=?", (key,)).fetchone()[0]
    m = dict(m, section=section, version=v)
    tid = c.execute("INSERT INTO templates(key,kind,section,name,version,status,mapping,created_by,created_at,sample_name,sample) VALUES(?,?,?,?,?,'draft',?,?,?,?,?)",
                    (key, kind, section, name or key, v, json.dumps(m), auth.current()["id"], A.now(), *sample)).lastrowid
    A.log(c, None, f"Template {key} version {v} drafted", kind="admin")
    return dict(id=tid, version=v)


def load_any(name, raw, series=None):
    """A data file as (content, kind, notes, templates): a workbook of logsheets is read with the templates; anything else
    (the section/field/value layout in CSV, Excel, JSON or SQLite) with importers.load_test_data as before."""
    if name.lower().endswith((".xlsx", ".xlsm")) or raw[:2] == b"PK":
        with A.db() as c: sheets, unmatched = read_workbook(c, name, raw)
        if sheets:
            errs = [f"{r['sheet']}: {e}" for r in sheets for e in r["errors"]]
            content, dup = to_content(sheets)
            if errs or dup: raise A.importers.ImportError_("; ".join(dup + errs))
            notes = [f"sheet '{r['sheet']}' read with template {r['template']['key']} v{r['template']['version']}" for r in sheets]
            if unmatched: notes.append(f"sheets not recognised and not read: {', '.join(unmatched)}")
            return content, "xlsx", notes, {r["section"]: r["template"]["id"] for r in sheets}
    content, kind, notes = A.importers.load_test_data(name, raw, series)
    return content, kind, notes, None


def install(app_module):
    global A
    A = app_module
    app, db, log, NAMES = A.app, A.db, A.log, A.NAMES
    with db() as c: init_db(c)

    # ---------------------------------------------------------------- registry (Admin manages, staff read)
    @app.get("/api/templates")
    @auth.require("staff.view")
    def templates_list():
        a = request.args; q, args = "SELECT * FROM templates WHERE 1=1", []
        for k in ("kind", "section", "status", "key"):
            if a.get(k): q += f" AND {k}=?"; args.append(a[k])
        with db() as c: return jsonify([row(r) for r in c.execute(q + " ORDER BY key, version DESC", args)])

    @app.get("/api/templates/<int:tid>")
    @auth.require("staff.view")
    def template_get(tid):
        with db() as c: return jsonify(row(template(c, tid), full=True))

    @app.post("/api/templates")
    @auth.require("templates.manage")
    def template_new():
        """New draft: a new version of an existing template (from_id), or a new template (key, kind, section)."""
        with db() as c:
            c.execute("BEGIN IMMEDIATE")
            res = create_template(c, A.body())
        return (jsonify(error=res), 409 if any("exists" in e for e in res) else 400) if isinstance(res, list) else (jsonify(res), 201)

    @app.post("/api/templates/<int:tid>")
    @auth.require("templates.manage")
    def template_edit(tid):
        b = A.body()
        with db() as c:
            r = template(c, tid)
            if r["status"] != "draft": return jsonify(error=["Only a draft can be changed; create a new version"]), 409
            m = b.get("mapping", json.loads(r["mapping"]))
            err = X.validate_mapping(m)
            if err: return jsonify(error=err), 400
            m = dict(m, section=r["section"], version=r["version"])
            c.execute("UPDATE templates SET mapping=?, name=? WHERE id=?", (json.dumps(m), str(b.get("name") or r["name"]), tid))
            log(c, None, f"Template {r['key']} version {r['version']} (draft) edited", kind="admin")
        return jsonify(ok=True)

    @app.post("/api/templates/<int:tid>/sample")
    @auth.require("templates.manage")
    def template_sample(tid):
        name, raw = file_arg(A.body()); book(name, raw)
        with db() as c:
            r = template(c, tid)
            if r["status"] != "draft": return jsonify(error=["Attach samples to a draft"]), 409
            c.execute("UPDATE templates SET sample=?, sample_name=? WHERE id=?", (raw, name, tid))
        return jsonify(ok=True)

    @app.post("/api/templates/grid")
    @auth.require("templates.manage")
    def template_grid():
        """The cells of a sample workbook, for clicking cells to bind them to fields."""
        b = A.body()
        if b.get("template_id"):
            with db() as c: r = template(c, b["template_id"])
            if r["sample"] is None: return jsonify(error=["This template has no sample file; upload one"]), 404
            name, raw = r["sample_name"], r["sample"]
        else: name, raw = file_arg(b)
        try: return jsonify(sheets=X.grid(raw, name))
        except X.TemplateError as e: return jsonify(error=[str(e)]), 400

    @app.post("/api/templates/test")
    @auth.require("templates.manage")
    def template_test():
        """Live preview: run a (possibly unsaved) mapping against a file or the template's sample."""
        b = A.body(); m = b.get("mapping")
        err = X.validate_mapping(m)
        if err: return jsonify(error=err), 400
        if b.get("template_id") and not b.get("file"):
            with db() as c: r = template(c, b["template_id"])
            if r["sample"] is None: return jsonify(error=["No sample file: upload one to test against"]), 400
            name, raw = r["sample_name"], r["sample"]
        else: name, raw = file_arg(b)
        bk = book(name, raw)
        hits = [ws for ws in bk.sheets() if X.matches(X.Sheet(bk, ws), m.get("fingerprint"))]
        try: res = X.extract(bk, m, hits[0] if hits else None)
        except X.TemplateError as e: return jsonify(error=[str(e)]), 400
        res["fingerprint_found"] = bool(hits)
        return jsonify(res)

    @app.post("/api/templates/<int:tid>/test-past")
    @auth.require("templates.manage")
    def template_test_past(tid):
        """Run the template over files already uploaded for its test and compare with what was stored from them."""
        n = min(int(A.body().get("limit") or 20), 100)
        with db() as c:
            t = template(c, tid); m = json.loads(t["mapping"])
            past = c.execute("""SELECT f.id, f.name, f.content, j.series, s.data FROM files f JOIN jobs j ON j.id=f.job_id
                                JOIN section_history s ON s.file_id=f.id AND s.key=? AND s.data IS NOT NULL
                                WHERE f.mime=? GROUP BY f.id ORDER BY f.id DESC LIMIT ?""", (t["section"], XLSX, n)).fetchall()
        out = []
        for p in past:
            try:
                bk = X.Book(p["content"], p["name"])
                hits = [ws for ws in bk.sheets() if X.matches(X.Sheet(bk, ws), m.get("fingerprint"))]
                if not hits: out.append(dict(file=p["name"], series=p["series"], result="not recognised")); continue
                r = X.extract(bk, m, hits[0])
                same = {k: v for k, v in r["data"].items() if v is not None} == {k: v for k, v in json.loads(p["data"]).items() if v is not None}
                out.append(dict(file=p["name"], series=p["series"], result="same as stored" if same and not r["errors"] else
                                "errors" if r["errors"] else "differs from stored", errors=r["errors"][:5]))
            except X.TemplateError as e: out.append(dict(file=p["name"], series=p["series"], result=f"unreadable: {e}"))
        return jsonify(files=out, checked=len(out))

    @app.post("/api/templates/<int:tid>/activate")
    @auth.require("templates.manage")
    def template_activate(tid):
        with db() as c:
            c.execute("BEGIN IMMEDIATE"); r = template(c, tid)
            if r["status"] != "draft": return jsonify(error=["Only a draft can be activated"]), 409
            err = X.validate_mapping(json.loads(r["mapping"]))
            if err: return jsonify(error=err), 400
            c.execute("UPDATE templates SET status='retired', retired_at=? WHERE key=? AND status='active'", (A.now(), r["key"]))
            c.execute("UPDATE templates SET status='active', activated_by=?, activated_at=? WHERE id=?", (auth.current()["id"], A.now(), tid))
            log(c, None, f"Template {r['key']} version {r['version']} activated (earlier version retired)", kind="admin")
        return jsonify(ok=True)

    @app.post("/api/templates/<int:tid>/retire")
    @auth.require("templates.manage")
    def template_retire(tid):
        with db() as c:
            r = template(c, tid)
            if r["status"] == "draft":
                c.execute("DELETE FROM templates WHERE id=?", (tid,)); log(c, None, f"Draft template {r['key']} v{r['version']} discarded", kind="admin")
            else:
                c.execute("UPDATE templates SET status='retired', retired_at=? WHERE id=?", (A.now(), tid))
                log(c, None, f"Template {r['key']} version {r['version']} retired (kept: reports made with it stay reproducible)", kind="admin")
        return jsonify(ok=True)

    @app.get("/api/templates/<int:a>/diff/<int:b>")
    @auth.require("staff.view")
    def template_diff(a, b):
        with db() as c: ma, mb = json.loads(template(c, a)["mapping"]), json.loads(template(c, b)["mapping"])
        return jsonify(X.diff(ma, mb))

    @app.get("/api/templates/<int:tid>/export")
    @auth.require("staff.view")
    def template_export(tid):
        with db() as c: r = template(c, tid)
        body = json.dumps(dict(key=r["key"], kind=r["kind"], section=r["section"], name=r["name"], version=r["version"], mapping=json.loads(r["mapping"])), indent=1)
        return send_file(io.BytesIO(body.encode()), mimetype="application/json", as_attachment=True, download_name=f"template_{r['key']}_v{r['version']}.json")

    @app.post("/api/templates/import")
    @auth.require("templates.manage")
    def template_import():
        """A mapping exported from another PC becomes a draft here (never active straight away)."""
        b = A.body(); t = b.get("template") if isinstance(b.get("template"), dict) else b
        with db() as c:
            c.execute("BEGIN IMMEDIATE")
            base = c.execute("SELECT id FROM templates WHERE key=? ORDER BY version DESC LIMIT 1", (t.get("key"),)).fetchone()
            res = create_template(c, dict(from_id=base[0], mapping=t.get("mapping"), name=t.get("name")) if base else
                                  dict(key=t.get("key"), kind=t.get("kind"), section=t.get("section"), name=t.get("name"), mapping=t.get("mapping")))
        return (jsonify(error=res), 400) if isinstance(res, list) else (jsonify(res), 201)

    # ---------------------------------------------------------------- blank sheets and forms
    @app.get("/api/logsheets/<section>.xlsx")
    @auth.require("staff.view")
    def logsheet_blank(section):
        """The current blank logsheet for one test, or every test in one workbook ('all')."""
        with db() as c: ts = active(c, "logsheet", None if section == "all" else section)
        if not ts: abort(404)
        order = list(NAMES)
        ts.sort(key=lambda t: order.index(t["section"]) if t["section"] in order else 99)
        raw = X.workbook([(t["mapping"], None, None) for t in ts])
        return send_file(io.BytesIO(raw), mimetype=XLSX, as_attachment=True,
                         download_name="Aletheia_logsheets_all.xlsx" if section == "all" else f"Aletheia_logsheet_{section}_v{ts[0]['version']}.xlsx")

    @app.get("/api/request-form.xlsx")
    @auth.require("jobs.view")
    def request_form_blank():
        """The customer request form as an Excel file (customers can download it from their portal too)."""
        with db() as c: ts = active(c, "request_form")
        if not ts: abort(404)
        return send_file(io.BytesIO(X.workbook([(ts[0]["mapping"], None, None)])), mimetype=XLSX, as_attachment=True,
                         download_name=f"CPRI_SCL_customer_request_form_v{ts[0]['version']}.xlsx")

    # ---------------------------------------------------------------- upload with preview (one job)
    @app.post("/api/jobs/<int:i>/excel/preview")
    @auth.require("data.write")
    def excel_preview(i):
        """What would be read from this workbook, value by value with its cell. Nothing is stored."""
        j = A.getjob(i, False); b = A.body(); name, raw = file_arg(b)
        with db() as c: sheets, unmatched = read_workbook(c, name, raw, b.get("template_id"))
        if b.get("section"):  # uploaded for one test: only that test's sheet is read
            A.test_key(b["section"]); unmatched += [r["sheet"] for r in sheets if r["section"] != b["section"]]
            sheets = [r for r in sheets if r["section"] == b["section"]]
        js = norm_series(j["series"])
        for r in sheets:
            s = norm_series(r["series_on_sheet"])
            r["series_matches"] = None if not s else (s == js or (len(s) >= 7 and js.endswith(s[-7:])))
        content, dup = to_content(sheets)
        return jsonify(sheets=sheets, unmatched=unmatched, sections=[k for k in content if k != "ids"], problems=dup,
                       ok=bool(sheets) and not dup and not any(r["errors"] for r in sheets))

    @app.post("/api/jobs/<int:i>/excel/import")
    @auth.require("data.write")
    def excel_import(i):
        """Store what the preview showed: the workbook is read again here (same file, same templates, same result)."""
        j = A.getjob(i); b = A.body(); name, raw = file_arg(b)
        if A.locked(j): return A.locked(j)
        bay = A.bay_of(b)
        with db() as c: sheets, _ = read_workbook(c, name, raw, b.get("template_id"))
        if b.get("sheets"): sheets = [r for r in sheets if r["sheet"] in b["sheets"]]
        if b.get("section"): A.test_key(b["section"]); sheets = [r for r in sheets if r["section"] == b["section"]]
        if not sheets: return jsonify(error=["No sheet of this workbook matches " + (f"the {A.NAMES.get(b['section'], 'chosen')} template" if b.get("section") else "a logsheet template")]), 400
        errs = [f"{r['sheet']}: {e}" for r in sheets for e in r["errors"]]
        content, dup = to_content(sheets)
        if errs or dup: return jsonify(error=dup + errs), 400
        notes = [f"sheet '{r['sheet']}' read with template {r['template']['key']} v{r['template']['version']}" for r in sheets]
        return A.apply_import(j, i, name, content, "xlsx", notes, raw, bay, templates={r["section"]: r["template"]["id"] for r in sheets})

    # ---------------------------------------------------------------- one workbook, several jobs
    @app.post("/api/excel/preview")
    @auth.require("data.write")
    def excel_multi_preview():
        """A workbook with sheets of several jobs: each sheet is routed by the series number written on it."""
        name, raw = file_arg(A.body())
        with db() as c:
            sheets, unmatched = read_workbook(c, name, raw)
            jobs = [dict(r) for r in c.execute("SELECT id, series, stage, archived FROM jobs WHERE archived=0")]
        for r in sheets:
            s = norm_series(r["series_on_sheet"]); r["job"] = None; r["route"] = "no series on the sheet"
            if s:
                exact = [x for x in jobs if norm_series(x["series"]) == s]
                near = [x for x in jobs if len(s) >= 7 and norm_series(x["series"]).endswith(s[-7:])]
                pick = exact or (near if len(near) == 1 else [])
                if pick: r["job"] = dict(id=pick[0]["id"], series=pick[0]["series"], released=pick[0]["stage"] == 4); r["route"] = "exact" if exact else "matched by the last 7 characters"
                else: r["route"] = "no open job with this series" if not near else f"{len(near)} jobs match; choose one"
        return jsonify(sheets=sheets, unmatched=unmatched)

    @app.post("/api/excel/import")
    @auth.require("data.write")
    def excel_multi_import():
        """Import the routed sheets: route = {sheet name: job id}. Each job gets its own copy of the file and its own audit entry."""
        b = A.body(); name, raw = file_arg(b); route = b.get("route") or {}
        bay = A.bay_of(b)
        with db() as c: sheets, _ = read_workbook(c, name, raw)
        per_job, errs = {}, []
        for r in sheets:
            jid = route.get(r["sheet"])
            if not jid: continue
            if r["errors"]: errs += [f"{r['sheet']}: {e}" for e in r["errors"]]
            per_job.setdefault(int(jid), []).append(r)
        if not per_job: return jsonify(error=["Choose a job for at least one sheet"]), 400
        for jid, rs in per_job.items():
            _, dup = to_content(rs); errs += dup
        if errs: return jsonify(error=errs), 400
        jobs = {jid: A.getjob(jid) for jid in per_job}  # every target exists and is visible, before anything is written
        results = {}
        for jid, rs in per_job.items():
            j = jobs[jid]
            if A.locked(j): results[j["series"]] = dict(status=409, error=A.locked(j)[0].get_json()["error"]); continue
            content, _ = to_content(rs)
            res = A.apply_import(j, jid, name, content, "xlsx", [f"sheet '{r['sheet']}' (template {r['template']['key']} v{r['template']['version']})" for r in rs],
                                 raw, bay, templates={r["section"]: r["template"]["id"] for r in rs})
            resp, code = (res, 200) if not isinstance(res, tuple) else res
            results[j["series"]] = dict(status=code, **resp.get_json())
        return jsonify(results=results)

    # ---------------------------------------------------------------- downloads across reports
    @app.get("/api/jobs/<int:i>/logsheets.xlsx")
    @auth.require("data.export")
    def job_logsheets(i):
        """The job's stored data written back into the logsheet layouts (the version that read each section, when known)."""
        j = A.getjob(i, False); d = j["data"]
        with db() as c:
            byid = {r["key"]: r["template_id"] for r in c.execute("SELECT key, template_id FROM sections WHERE job_id=?", (i,))}
            cur = {t["section"]: t for t in active(c, "logsheet")}
            sheets = []
            for k in NAMES:
                if k not in d or k == "request": continue
                t = None
                if byid.get(k):
                    r = c.execute("SELECT * FROM templates WHERE id=?", (byid[k],)).fetchone()
                    if r: t = dict(mapping=json.loads(r["mapping"]))
                t = t or cur.get(k)
                if t: sheets.append((t["mapping"], d[k], d.get("ids") or {}))
        if not sheets: return jsonify(error=["No logsheet data to export"]), 409
        return send_file(io.BytesIO(X.workbook(sheets)), mimetype=XLSX, as_attachment=True, download_name=f"{j['series']}_logsheets.xlsx")

    @app.get("/api/records.xlsx")
    @auth.require("data.export")
    def records_xlsx():
        """Many test reports in one workbook: one row per job, per-test status, and key logged values side by side."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        ids = [int(x) for x in request.args.get("ids", "").split(",") if x.strip().isdigit()]
        sel = ids or [r["id"] for r in A.jobs().get_json()]  # the same filters as the records list (q, stage, verdict, dates)
        wb = Workbook(); ws = wb.active; ws.title = "Jobs"
        head = ["Series", "Sample", "Customer", "Rating", "Status", "Result", "Received", "Approved by", "Employee ID", "Report version", "Report SHA-256"]
        ws.append(head)
        tests = wb.create_sheet("Tests"); keys = [k for k in NAMES if k != "request"]
        tests.append(["Series"] + [NAMES[k] for k in keys])
        vals = wb.create_sheet("Key values (as logged)")
        KEY = [("proforma", "kva", "Rating kVA"), ("proforma", "z_pct", "Declared %Z"), ("proforma", "loss100", "Guaranteed loss 100% W"),
               ("losses", "nll_at", "No-load loss after SC W"), ("temp", "oil_rise_reported", "Top-oil rise K (logged)"),
               ("temp", "hv_rise", "HV winding rise K (logged)"), ("temp", "lv_rise", "LV winding rise K (logged)"),
               ("noload", "v100", "Rated voltage V"), ("sc", "date", "SC test date"), ("pressure", "leak.obs", "Oil leakage observation")]
        vals.append(["Series"] + [l for _, _, l in KEY])
        for jid in sel:
            try: j = A.getjob(jid)
            except Exception: continue  # noqa: BLE001 - not visible or gone
            rel = next((r for r in j["reports"] if r["approver"]), None)
            ws.append([j["series"], j["sample"], j["customer"], j["rating"], j["stage_name"], j.get("verdict") or "", (j.get("created") or "")[:10],
                       j.get("approver") or "", j.get("approver_id") or "", rel["version"] if rel else "", rel["sha256"] if rel else ""])
            st = {p["key"]: p["state"] for p in j["progress"]}
            tests.append([j["series"]] + [{"not_started": "", "uploaded": "awaiting verification"}.get(st.get(k, (j["meta"].get(k) or {}).get("state", "")), st.get(k, (j["meta"].get(k) or {}).get("state", ""))) for k in keys])
            vals.append([j["series"]] + [X.get(j["data"].get(s) or {}, p) if (j["meta"].get(s) or {}).get("state") in ("uploaded", "verified") else None for s, p, _ in KEY])
        for sh in (ws, tests, vals):
            for c in sh[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DFE7F1")
            for col in sh.columns: sh.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col) + 2))
            sh.freeze_panes = "B2"
        out = io.BytesIO(); wb.save(out)
        with db() as c: log(c, None, f"Records workbook exported ({len(sel)} jobs)", kind="event")
        return send_file(io.BytesIO(out.getvalue()), mimetype=XLSX, as_attachment=True, download_name=f"Aletheia_records_{dt.date.today():%Y%m%d}.xlsx")
